"""PDF y Excel de los ajustes de inventario: el reporte del período y el
documento de un ajuste (AJ-000001) para imprimir y firmar."""
from io import BytesIO
from xml.sax.saxutils import escape

from app.utils.calculations import to_decimal
from app.utils.turnos_caja import a_hora_local

AZUL = "#1d4ed8"


def _estilos():
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    base = getSampleStyleSheet()
    return {
        "titulo": ParagraphStyle("t", parent=base["Title"], fontSize=15, alignment=0, spaceAfter=2),
        "sub": ParagraphStyle("s", parent=base["Normal"], fontSize=9.5, textColor=colors.HexColor("#475569")),
        "seccion": ParagraphStyle("h", parent=base["Normal"], fontSize=10.5, fontName="Helvetica-Bold",
                                  textColor=colors.HexColor(AZUL), spaceBefore=10, spaceAfter=4),
        "celda": ParagraphStyle("c", parent=base["Normal"], fontSize=7.8, leading=9.5),
    }


def _tabla(datos, anchos, num_cols=(), negrita_ultima=False):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    t = Table(datos, colWidths=anchos, repeatRows=1)
    estilo = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(AZUL)), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.8),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
    ]
    for c in num_cols:
        estilo.append(("ALIGN", (c, 0), (c, -1), "RIGHT"))
    if negrita_ultima:
        estilo.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
    t.setStyle(TableStyle(estilo))
    t.hAlign = "LEFT"
    return t


def _resumen(pares, moneda):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    t = Table([[p[0] for p in pares], [p[1] for p in pares]])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, 0), 8), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#475569")),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"), ("FONTSIZE", (0, 1), (-1, 1), 12),
        ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 20),
    ]))
    t.hAlign = "LEFT"
    return t


def _m(v, moneda):
    return f"{moneda} {to_decimal(v or 0):,.2f}"


# ---------------------------------------------------------------------------
# Reporte del período
# ---------------------------------------------------------------------------
def pdf_reporte_ajustes(datos: dict, rango_txt: str, nombre_taller: str, moneda: str, filtros: str = "") -> bytes:
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    e = _estilos()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=1.1 * cm, bottomMargin=1.1 * cm,
                            leftMargin=1.2 * cm, rightMargin=1.2 * cm, title="Reporte de ajustes de inventario")
    t = datos["total"]
    story = [Paragraph(escape(nombre_taller), e["titulo"]),
             Paragraph(f"Reporte de ajustes de inventario · {escape(rango_txt)}{(' · ' + escape(filtros)) if filtros else ''}", e["sub"]),
             Spacer(1, 8),
             _resumen([("Unidades que subieron", f"{t['unidades_entrada']}"), ("Unidades que bajaron", f"{t['unidades_salida']}"),
                       ("Valor que subió", _m(t["valor_entrada"], moneda)), ("Valor que bajó", _m(t["valor_salida"], moneda)),
                       ("Neto", _m(t["valor_neto"], moneda))], moneda)]

    def grupo(titulo, filas, nombre_col):
        if not filas:
            return
        story.append(Paragraph(titulo, e["seccion"]))
        datos_t = [[nombre_col, "Suben (u.)", "Bajan (u.)", "Valor que subió", "Valor que bajó", "Neto"]]
        for g in filas:
            datos_t.append([Paragraph(escape(str(g["clave"])), e["celda"]), g["unidades_entrada"], g["unidades_salida"],
                            _m(g["valor_entrada"], moneda), _m(g["valor_salida"], moneda), _m(g["valor_neto"], moneda)])
        story.append(_tabla(datos_t, [8.5 * cm, 2.4 * cm, 2.4 * cm, 3.4 * cm, 3.4 * cm, 3.4 * cm], num_cols=(1, 2, 3, 4, 5)))

    grupo("Por motivo", datos["por_motivo"], "Motivo")
    grupo("Por categoría", datos["por_categoria"], "Categoría")
    grupo("Por mes", datos["por_mes"], "Mes")
    if datos["por_producto"]:
        story.append(Paragraph("Por producto", e["seccion"]))
        datos_t = [["Producto", "Suben (u.)", "Bajan (u.)", "Neto (u.)", "Valor que subió", "Valor que bajó", "Neto"]]
        for g in datos["por_producto"]:
            datos_t.append([Paragraph(escape(str(g["clave"])), e["celda"]), g["unidades_entrada"], g["unidades_salida"], g["unidades_neto"],
                            _m(g["valor_entrada"], moneda), _m(g["valor_salida"], moneda), _m(g["valor_neto"], moneda)])
        story.append(_tabla(datos_t, [8.1 * cm, 2.1 * cm, 2.1 * cm, 2.1 * cm, 3.1 * cm, 3.1 * cm, 3.1 * cm], num_cols=(1, 2, 3, 4, 5, 6)))

    story.append(Paragraph("Detalle", e["seccion"]))
    datos_t = [["Fecha", "Ajuste", "Motivo", "Producto", "Dif.", "Costo", "Valor", "Usuario"]]
    for f in datos["filas"]:
        producto = f"{f['codigo']} · {f['nombre']}" + (f" ({f['categoria']})" if f["categoria"] else "")
        datos_t.append([a_hora_local(f["fecha"]).strftime("%d/%m/%Y %I:%M %p"), f["ajuste"].numero if f["ajuste"] else "—",
                        Paragraph(escape(f["motivo"] + (f" · {f['detalle_motivo']}" if f.get("detalle_motivo") else "")), e["celda"]),
                        Paragraph(escape(producto), e["celda"]), f"{f['diferencia']:+d}",
                        _m(f["costo_unitario"], moneda) + (" *" if f["aproximado"] else ""), _m(f["valor"], moneda), escape(f["usuario"] or "")])
    if len(datos_t) == 1:
        datos_t.append(["", "", "Sin ajustes en el período.", "", "", "", "", ""])
    story.append(_tabla(datos_t, [3.2 * cm, 2.2 * cm, 4.8 * cm, 6.6 * cm, 1.3 * cm, 2.4 * cm, 2.5 * cm, 2.3 * cm], num_cols=(4, 5, 6)))
    if datos.get("hay_aproximados"):
        story += [Spacer(1, 6), Paragraph("* Ajustes rápidos anteriores al módulo: valorados con el costo actual del producto (aproximado).", e["sub"])]
    doc.build(story)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# Documento de un ajuste
# ---------------------------------------------------------------------------
def pdf_ajuste(ajuste, nombre_taller: str, moneda: str, revierte_a=None, revertido_por=None) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    e = _estilos()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=1.3 * cm, bottomMargin=1.3 * cm,
                            leftMargin=1.4 * cm, rightMargin=1.4 * cm, title=f"Ajuste {ajuste.numero}")
    origen = {"manual": "en pantalla", "excel": "desde Excel", "reversion": "reversión"}.get(ajuste.origen, ajuste.origen)
    story = [Paragraph(escape(nombre_taller), e["titulo"]),
             Paragraph(f"Ajuste de inventario <b>{escape(ajuste.numero)}</b> · {a_hora_local(ajuste.fecha).strftime('%d/%m/%Y %I:%M %p')}", e["sub"]),
             Paragraph(f"Motivo: <b>{escape(ajuste.motivo)}</b> · Hecho por {escape(ajuste.usuario_nombre or '')} ({origen})", e["sub"])]
    if ajuste.observacion:
        story.append(Paragraph(f"Observación: {escape(ajuste.observacion)}", e["sub"]))
    if revierte_a:
        story.append(Paragraph(f"Este ajuste deshace el {escape(revierte_a.numero)}.", e["sub"]))
    if revertido_por:
        story.append(Paragraph(f"REVERTIDO por el ajuste {escape(revertido_por.numero)}.", e["sub"]))
    story += [Spacer(1, 8), _resumen([("Subieron", f"{ajuste.unidades_entrada} u. · {_m(ajuste.valor_entrada, moneda)}"),
                                       ("Bajaron", f"{ajuste.unidades_salida} u. · {_m(ajuste.valor_salida, moneda)}"),
                                       ("Valor neto (al costo)", _m(ajuste.valor_neto, moneda))], moneda),
              Spacer(1, 10)]
    datos_t = [["Código", "Producto", "Había", "Quedó", "Dif.", "Costo", "Valor"]]
    for d in ajuste.detalles:
        datos_t.append([d.producto_codigo, Paragraph(escape(d.producto_nombre or ""), e["celda"]), d.existencia_anterior, d.existencia_nueva,
                        f"{d.diferencia:+d}", _m(d.costo_unitario, moneda), _m(d.valor, moneda)])
    datos_t.append(["", "Total", "", "", f"{sum(d.diferencia for d in ajuste.detalles):+d}", "", _m(ajuste.valor_neto, moneda)])
    story.append(_tabla(datos_t, [2.6 * cm, 6.6 * cm, 1.6 * cm, 1.6 * cm, 1.5 * cm, 2.4 * cm, 2.5 * cm], num_cols=(2, 3, 4, 5, 6), negrita_ultima=True))
    firma = Table([["_____________________________", "_____________________________"],
                   ["Hecho por", "Revisado por"]], colWidths=[8.5 * cm, 8.5 * cm])
    firma.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, 0), 40)]))
    story += [Spacer(1, 10), firma]
    doc.build(story)
    return buffer.getvalue()


def excel_ajuste(ajuste, nombre_taller: str, moneda: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    h = wb.active
    h.title = ajuste.numero
    h["A1"] = nombre_taller
    h["A1"].font = Font(size=14, bold=True)
    h["A2"] = f"Ajuste de inventario {ajuste.numero} · {a_hora_local(ajuste.fecha).strftime('%d/%m/%Y %I:%M %p')}"
    h["A2"].font = Font(size=12, bold=True)
    h["A3"] = f"Motivo: {ajuste.motivo} · Hecho por {ajuste.usuario_nombre or ''}"
    if ajuste.observacion:
        h["A4"] = f"Observación: {ajuste.observacion}"
    fila = 6
    for col, t in enumerate(["Código", "Producto", "Había", "Quedó", "Diferencia", f"Costo ({moneda})", f"Valor ({moneda})"], start=1):
        c = h.cell(row=fila, column=col, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    for d in ajuste.detalles:
        fila += 1
        for col, v in enumerate([d.producto_codigo, d.producto_nombre, d.existencia_anterior, d.existencia_nueva, d.diferencia,
                                 float(d.costo_unitario or 0), float(d.valor or 0)], start=1):
            c = h.cell(row=fila, column=col, value=v)
            if col >= 6:
                c.number_format = '#,##0.00'
    fila += 1
    h.cell(row=fila, column=2, value="Total").font = Font(bold=True)
    h.cell(row=fila, column=5, value=sum(d.diferencia for d in ajuste.detalles)).font = Font(bold=True)
    c = h.cell(row=fila, column=7, value=float(ajuste.valor_neto or 0))
    c.font = Font(bold=True)
    c.number_format = '#,##0.00'
    for col, ancho in zip("ABCDEFG", (16, 44, 10, 10, 12, 14, 14)):
        h.column_dimensions[col].width = ancho
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
