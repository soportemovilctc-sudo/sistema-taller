"""Estados Financieros (solo administrador): Estado de Resultados, Flujo de
Efectivo, Balance General y Cambios en el Patrimonio, por mes o por rango de
fechas, en pantalla y exportables a Excel y PDF. Más "Datos del balance",
donde el admin registra lo que el sistema no sabe por sí solo (activos
fijos, préstamos, deudas, capital, retiros). Los cálculos están en
app/utils/estados_financieros.py."""
import io
from xml.sax.saxutils import escape
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from fastapi import APIRouter, Request, Depends, Form
from fastapi.responses import RedirectResponse, StreamingResponse
from sqlalchemy.orm import Session

from app.templates_env import templates
from app.database import get_db
from app.deps import roles_required
from app.models import Configuracion, PartidaBalance, TIPOS_PARTIDA_BALANCE
from app.utils.flash import flash
from app.utils.calculations import to_decimal
from app.utils.turnos_caja import hoy_local
from app.utils.estados_financieros import Periodo, periodo_mes, generar_estados, NOMBRES_MES

router = APIRouter()
BASE = "/estados-financieros"
solo_admin = roles_required("admin")
VISTAS = {
    "resultados": "Estado de Resultados",
    "flujo": "Flujo de Efectivo",
    "balance": "Balance General",
    "patrimonio": "Cambios en el Patrimonio",
}


def _fecha(valor: str):
    try:
        return datetime.strptime((valor or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _periodo_desde_query(modo: str, anio: int, mes: int, desde: str, hasta: str) -> tuple[Periodo, dict]:
    hoy = hoy_local()
    if modo == "rango":
        d, h = _fecha(desde), _fecha(hasta)
        if d and h:
            periodo = Periodo(d, h)
            return periodo, {"modo": "rango", "anio": hoy.year, "mes": hoy.month,
                             "desde": periodo.desde.isoformat(), "hasta": periodo.hasta.isoformat()}
    anio = anio if 2000 <= (anio or 0) <= 2100 else hoy.year
    mes = mes if 1 <= (mes or 0) <= 12 else hoy.month
    periodo = periodo_mes(anio, mes)
    return periodo, {"modo": "mes", "anio": anio, "mes": mes,
                     "desde": periodo.desde.isoformat(), "hasta": periodo.hasta.isoformat()}


def _nombre_taller(db: Session) -> str:
    cfg = db.get(Configuracion, 1)
    return cfg.nombre_taller if cfg and cfg.nombre_taller else "Mi Taller"


def _query_string(sel: dict) -> str:
    if sel["modo"] == "rango":
        return f"modo=rango&desde={sel['desde']}&hasta={sel['hasta']}"
    return f"modo=mes&anio={sel['anio']}&mes={sel['mes']}"


# ---------------------------------------------------------------------------
# Pantalla
# ---------------------------------------------------------------------------
@router.get(BASE)
def estados_index(
    request: Request, vista: str = "resultados", modo: str = "mes", anio: int = 0, mes: int = 0,
    desde: str = "", hasta: str = "",
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    if vista not in VISTAS:
        vista = "resultados"
    periodo, sel = _periodo_desde_query(modo, anio, mes, desde, hasta)
    estados = generar_estados(db, periodo)
    hoy = hoy_local()
    return templates.TemplateResponse("estados_financieros/index.html", {
        "request": request, "usuario": usuario, "vista": vista, "vistas": VISTAS,
        "estado": estados[vista], "estados": estados, "sel": sel, "qs": _query_string(sel),
        "meses": [(i, NOMBRES_MES[i].capitalize()) for i in range(1, 13)],
        "anios": list(range(hoy.year - 5, hoy.year + 1)),
        "hoy_iso": hoy.isoformat(), "nombre_taller": _nombre_taller(db),
    })


# ---------------------------------------------------------------------------
# Exportar a Excel (los 4 estados, una hoja cada uno)
# ---------------------------------------------------------------------------
@router.get(BASE + "/exportar/excel")
def estados_excel(
    modo: str = "mes", anio: int = 0, mes: int = 0, desde: str = "", hasta: str = "",
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    periodo, sel = _periodo_desde_query(modo, anio, mes, desde, hasta)
    estados = generar_estados(db, periodo)
    taller = _nombre_taller(db)
    wb = Workbook()
    wb.remove(wb.active)
    relleno_titulo = PatternFill("solid", fgColor="1F4E78")
    linea_total = Border(top=Side(style="thin"), bottom=Side(style="double"))
    linea_sub = Border(top=Side(style="thin"))
    formato = '#,##0.00;[Red]-#,##0.00'

    for clave, nombre_hoja in (("resultados", "Estado de Resultados"), ("flujo", "Flujo de Efectivo"),
                               ("balance", "Balance General"), ("patrimonio", "Cambios Patrimonio")):
        e = estados[clave]
        hoja = wb.create_sheet(nombre_hoja)
        hoja.column_dimensions["A"].width = 72
        hoja.column_dimensions["B"].width = 20
        hoja.append([taller])
        hoja["A1"].font = Font(bold=True, size=14)
        hoja.append([e["titulo"]])
        hoja["A2"].font = Font(bold=True, size=12)
        hoja.append([e["periodo"] + " · cifras en lempiras"])
        hoja.append([])
        for linea in e["lineas"]:
            sangria = "    " * linea["nivel"]
            hoja.append([sangria + linea["etiqueta"], float(linea["monto"]) if linea["monto"] is not None and linea["tipo"] != "titulo" else None])
            fila = hoja.max_row
            celda_a, celda_b = hoja.cell(row=fila, column=1), hoja.cell(row=fila, column=2)
            celda_b.number_format = formato
            celda_b.alignment = Alignment(horizontal="right")
            if linea["tipo"] == "titulo" and linea["nivel"] == 0:
                celda_a.font = Font(bold=True, color="FFFFFF")
                celda_a.fill = relleno_titulo
                celda_b.fill = relleno_titulo
            elif linea["tipo"] == "titulo":
                celda_a.font = Font(bold=True)
            elif linea["tipo"] == "subtotal":
                celda_a.font = celda_b.font = Font(bold=True)
                celda_b.border = linea_sub
            elif linea["tipo"] == "total":
                celda_a.font = celda_b.font = Font(bold=True, size=12)
                celda_b.border = linea_total
        if e["notas"]:
            hoja.append([])
            hoja.append(["Notas"])
            hoja.cell(row=hoja.max_row, column=1).font = Font(bold=True)
            for n in e["notas"]:
                hoja.append(["• " + n])

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    nombre = f"estados_financieros_{sel['desde']}_{sel['hasta']}.xlsx"
    return StreamingResponse(salida, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


# ---------------------------------------------------------------------------
# Exportar a PDF (los 4 estados, uno por página)
# ---------------------------------------------------------------------------
@router.get(BASE + "/exportar/pdf")
def estados_pdf(
    modo: str = "mes", anio: int = 0, mes: int = 0, desde: str = "", hasta: str = "",
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    from reportlab.lib.pagesizes import letter
    from reportlab.lib import colors
    from reportlab.lib.units import cm
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak

    periodo, sel = _periodo_desde_query(modo, anio, mes, desde, hasta)
    estados = generar_estados(db, periodo)
    taller = _nombre_taller(db)
    estilos = getSampleStyleSheet()
    est_taller = ParagraphStyle("taller", parent=estilos["Title"], fontSize=15, spaceAfter=2)
    est_titulo = ParagraphStyle("titulo", parent=estilos["Heading2"], alignment=1, spaceAfter=0)
    est_periodo = ParagraphStyle("periodo", parent=estilos["Normal"], alignment=1, fontSize=9, textColor=colors.grey)
    est_celda = ParagraphStyle("celda", parent=estilos["Normal"], fontSize=9, leading=11)
    est_nota = ParagraphStyle("nota", parent=estilos["Normal"], fontSize=8, textColor=colors.grey, leading=10)

    def monto(v):
        if v is None:
            return ""
        v = Decimal(v)
        return f"({abs(v):,.2f})" if v < 0 else f"{v:,.2f}"

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm,
                            title=f"Estados financieros {periodo.texto}")
    story = []
    for i, clave in enumerate(("resultados", "flujo", "balance", "patrimonio")):
        e = estados[clave]
        if i:
            story.append(PageBreak())
        story += [Paragraph(escape(taller), est_taller), Paragraph(e["titulo"], est_titulo),
                  Paragraph(f"{e['periodo']} · cifras en lempiras", est_periodo), Spacer(1, 0.4 * cm)]
        datos, estilos_tabla = [], [("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (1, 0), (1, -1), "RIGHT"),
                                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2), ("TOPPADDING", (0, 0), (-1, -1), 2)]
        for fila, linea in enumerate(e["lineas"]):
            sangria = "&nbsp;" * (4 * linea["nivel"])
            texto = sangria + escape(linea["etiqueta"])
            if linea["tipo"] in ("titulo", "subtotal", "total"):
                texto = f"<b>{texto}</b>"
            datos.append([Paragraph(texto, est_celda), "" if linea["tipo"] == "titulo" else monto(linea["monto"])])
            if linea["tipo"] == "titulo" and linea["nivel"] == 0:
                estilos_tabla += [("BACKGROUND", (0, fila), (-1, fila), colors.HexColor("#1F4E78")),
                                  ("TEXTCOLOR", (0, fila), (-1, fila), colors.white)]
                datos[-1][0] = Paragraph(f"<font color='white'><b>{escape(linea['etiqueta'])}</b></font>", est_celda)
            elif linea["tipo"] == "subtotal":
                estilos_tabla += [("LINEABOVE", (1, fila), (1, fila), 0.5, colors.black),
                                  ("FONTNAME", (1, fila), (1, fila), "Helvetica-Bold")]
            elif linea["tipo"] == "total":
                estilos_tabla += [("LINEABOVE", (1, fila), (1, fila), 0.8, colors.black),
                                  ("LINEBELOW", (1, fila), (1, fila), 1.2, colors.black),
                                  ("FONTNAME", (1, fila), (1, fila), "Helvetica-Bold"),
                                  ("BACKGROUND", (0, fila), (-1, fila), colors.HexColor("#EEF3F9"))]
        tabla = Table(datos, colWidths=[12.6 * cm, 4.4 * cm])
        tabla.setStyle(TableStyle(estilos_tabla))
        story.append(tabla)
        if e["notas"]:
            story.append(Spacer(1, 0.4 * cm))
            for n in e["notas"]:
                story.append(Paragraph("• " + escape(n), est_nota))
        story.append(Spacer(1, 0.3 * cm))
        story.append(Paragraph(f"Generado el {hoy_local().strftime('%d/%m/%Y')} por {escape(usuario['nombre_completo'])}. "
                               "Las cifras negativas se muestran entre paréntesis.", est_nota))
    doc.build(story)
    buffer.seek(0)
    nombre = f"estados_financieros_{sel['desde']}_{sel['hasta']}.pdf"
    return StreamingResponse(buffer, media_type="application/pdf",
                             headers={"Content-Disposition": f'attachment; filename="{nombre}"'})


# ---------------------------------------------------------------------------
# Datos del balance (activos fijos, deudas, préstamos, capital, retiros)
# ---------------------------------------------------------------------------
URL_DATOS = BASE + "/datos-balance"


def _leer_partida(tipo, descripcion, monto, fecha, fecha_baja, vida_util_meses):
    if tipo not in TIPOS_PARTIDA_BALANCE:
        raise ValueError("Elige el tipo.")
    if not descripcion.strip():
        raise ValueError("Escribe una descripción.")
    try:
        monto_dec = to_decimal(monto)
    except (InvalidOperation, ValueError):
        raise ValueError("El monto debe ser un número.")
    if monto_dec <= 0:
        raise ValueError("El monto debe ser mayor que cero.")
    f = _fecha(fecha)
    if not f:
        raise ValueError("Elige la fecha.")
    fb = _fecha(fecha_baja) if fecha_baja.strip() else None
    if fecha_baja.strip() and not fb:
        raise ValueError("La fecha de baja no es válida.")
    if fb and fb <= f:
        raise ValueError("La fecha de baja debe ser posterior a la fecha de registro.")
    try:
        vida = int(vida_util_meses or 0)
    except ValueError:
        raise ValueError("La vida útil debe ser un número entero de meses.")
    if vida < 0:
        raise ValueError("La vida útil no puede ser negativa.")
    if tipo != "activo_fijo":
        vida = 0
    return tipo, descripcion.strip(), monto_dec.quantize(Decimal("0.01")), f, fb, vida


@router.get(URL_DATOS)
def datos_balance(request: Request, editar: int = 0, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    partidas = db.query(PartidaBalance).order_by(PartidaBalance.tipo, PartidaBalance.fecha).all()
    grupos = []
    for clave, nombre in TIPOS_PARTIDA_BALANCE.items():
        filas = [p for p in partidas if p.tipo == clave]
        if filas:
            grupos.append({"clave": clave, "nombre": nombre, "filas": filas,
                           "total": sum((to_decimal(p.monto) for p in filas if p.fecha_baja is None), Decimal("0.00"))})
    return templates.TemplateResponse("estados_financieros/datos_balance.html", {
        "request": request, "usuario": usuario, "grupos": grupos, "tipos": TIPOS_PARTIDA_BALANCE,
        "en_edicion": db.get(PartidaBalance, editar) if editar else None,
        "hoy_iso": hoy_local().isoformat(),
    })


@router.post(URL_DATOS)
def datos_balance_crear(
    request: Request, tipo: str = Form(""), descripcion: str = Form(""), monto: str = Form("0"),
    fecha: str = Form(""), fecha_baja: str = Form(""), vida_util_meses: str = Form("0"), notas: str = Form(""),
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    try:
        tipo, descripcion, monto_dec, f, fb, vida = _leer_partida(tipo, descripcion, monto, fecha, fecha_baja, vida_util_meses)
    except ValueError as e:
        flash(request, str(e), "error")
        return RedirectResponse(URL_DATOS, status_code=303)
    db.add(PartidaBalance(tipo=tipo, descripcion=descripcion, monto=monto_dec, fecha=f, fecha_baja=fb,
                          vida_util_meses=vida, notas=notas.strip(), usuario_nombre=usuario["nombre_completo"],
                          creado_en=datetime.utcnow(), actualizado_en=datetime.utcnow()))
    db.commit()
    flash(request, f"Registrado: {descripcion}.", "success")
    return RedirectResponse(URL_DATOS, status_code=303)


@router.post(URL_DATOS + "/{partida_id}/editar")
def datos_balance_editar(
    partida_id: int, request: Request, tipo: str = Form(""), descripcion: str = Form(""), monto: str = Form("0"),
    fecha: str = Form(""), fecha_baja: str = Form(""), vida_util_meses: str = Form("0"), notas: str = Form(""),
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    partida = db.get(PartidaBalance, partida_id)
    if not partida:
        flash(request, "Registro no encontrado.", "error")
        return RedirectResponse(URL_DATOS, status_code=303)
    try:
        tipo, descripcion, monto_dec, f, fb, vida = _leer_partida(tipo, descripcion, monto, fecha, fecha_baja, vida_util_meses)
    except ValueError as e:
        flash(request, str(e), "error")
        return RedirectResponse(f"{URL_DATOS}?editar={partida_id}", status_code=303)
    partida.tipo, partida.descripcion, partida.monto = tipo, descripcion, monto_dec
    partida.fecha, partida.fecha_baja, partida.vida_util_meses = f, fb, vida
    partida.notas = notas.strip()
    partida.usuario_nombre = usuario["nombre_completo"]
    partida.actualizado_en = datetime.utcnow()
    db.commit()
    flash(request, f"Actualizado: {descripcion}.", "success")
    return RedirectResponse(URL_DATOS, status_code=303)


@router.post(URL_DATOS + "/{partida_id}/eliminar")
def datos_balance_eliminar(partida_id: int, request: Request, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    partida = db.get(PartidaBalance, partida_id)
    if partida:
        db.delete(partida)
        db.commit()
        flash(request, "Registro eliminado.", "success")
    return RedirectResponse(URL_DATOS, status_code=303)
