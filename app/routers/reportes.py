"""Módulo de reportes con filtros y exportación a CSV y Excel."""
import csv
import io
from datetime import datetime, date

from fastapi import APIRouter, Request, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.templates_env import templates
from app.database import get_db
from app.models import (
    OrdenServicio, MovimientoFinanciero, Venta, Producto, Tecnico,
    ESTADOS_ORDEN, FORMAS_PAGO,
)
from app.deps import login_required

router = APIRouter()


def _parse_fecha(valor, default=None):
    if not valor:
        return default
    try:
        return datetime.strptime(valor, "%Y-%m-%d")
    except ValueError:
        return default


def _filtrar_ordenes(db, desde, hasta, tecnico_id, estado):
    query = db.query(OrdenServicio).options(joinedload(OrdenServicio.cliente), joinedload(OrdenServicio.tecnico))
    if desde:
        query = query.filter(OrdenServicio.fecha >= desde)
    if hasta:
        query = query.filter(OrdenServicio.fecha <= hasta)
    if tecnico_id:
        query = query.filter(OrdenServicio.tecnico_id == int(tecnico_id))
    if estado:
        query = query.filter(OrdenServicio.estado == estado)
    return query.order_by(OrdenServicio.fecha.desc()).all()


@router.get("/reportes")
def reportes_index(
    request: Request, tipo: str = "ordenes",
    desde: str = "", hasta: str = "", tecnico_id: str = "", estado: str = "", forma_pago: str = "",
    db: Session = Depends(get_db), usuario=Depends(login_required),
):
    fecha_desde = _parse_fecha(desde)
    fecha_hasta = _parse_fecha(hasta)
    if fecha_hasta:
        fecha_hasta = fecha_hasta.replace(hour=23, minute=59, second=59)

    contexto = {
        "request": request, "usuario": usuario, "tipo": tipo,
        "desde": desde, "hasta": hasta, "tecnico_id": tecnico_id, "estado": estado, "forma_pago": forma_pago,
        "tecnicos": db.query(Tecnico).order_by(Tecnico.nombre).all(),
        "estados": ESTADOS_ORDEN, "formas_pago": FORMAS_PAGO,
    }

    if tipo == "ordenes" or tipo == "reparaciones":
        contexto["ordenes"] = _filtrar_ordenes(db, fecha_desde, fecha_hasta, tecnico_id, estado)
    elif tipo == "ingresos":
        q = db.query(MovimientoFinanciero).filter(MovimientoFinanciero.tipo == "ingreso")
        if fecha_desde:
            q = q.filter(MovimientoFinanciero.fecha >= fecha_desde)
        if fecha_hasta:
            q = q.filter(MovimientoFinanciero.fecha <= fecha_hasta)
        contexto["movimientos"] = q.order_by(MovimientoFinanciero.fecha.desc()).all()
    elif tipo == "gastos":
        q = db.query(MovimientoFinanciero).filter(MovimientoFinanciero.tipo == "gasto")
        if fecha_desde:
            q = q.filter(MovimientoFinanciero.fecha >= fecha_desde)
        if fecha_hasta:
            q = q.filter(MovimientoFinanciero.fecha <= fecha_hasta)
        contexto["movimientos"] = q.order_by(MovimientoFinanciero.fecha.desc()).all()
    elif tipo == "ventas":
        q = db.query(Venta)
        if fecha_desde:
            q = q.filter(Venta.fecha >= fecha_desde)
        if fecha_hasta:
            q = q.filter(Venta.fecha <= fecha_hasta)
        if forma_pago:
            q = q.filter(Venta.forma_pago == forma_pago)
        contexto["ventas"] = q.order_by(Venta.fecha.desc()).all()
    elif tipo == "inventario":
        contexto["productos"] = db.query(Producto).order_by(Producto.nombre).all()
    elif tipo == "saldos":
        contexto["ordenes"] = db.query(OrdenServicio).options(joinedload(OrdenServicio.cliente)).filter(
            OrdenServicio.saldo > 0, OrdenServicio.estado != "CANCELADO"
        ).order_by(OrdenServicio.fecha.desc()).all()

    return templates.TemplateResponse("reportes/index.html", contexto)


def _csv_response(filas, encabezados, nombre_archivo):
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(encabezados)
    writer.writerows(filas)
    buffer.seek(0)
    return StreamingResponse(iter([buffer.getvalue()]), media_type="text/csv",
                              headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'})


def _xlsx_response(filas, encabezados, nombre_archivo, titulo_hoja="Reporte"):
    """Genera un Excel (.xlsx) con encabezado resaltado, columnas
    ajustadas al contenido, formato de número en las columnas numéricas y
    la fila de encabezado congelada, con el mismo estilo en todos los
    reportes del sistema."""
    wb = Workbook()
    hoja = wb.active
    hoja.title = titulo_hoja[:31] or "Reporte"

    encabezado_fill = PatternFill(start_color="121D33", end_color="121D33", fill_type="solid")
    encabezado_font = Font(bold=True, color="FFFFFF")
    for col, titulo in enumerate(encabezados, start=1):
        celda = hoja.cell(row=1, column=col, value=titulo)
        celda.fill = encabezado_fill
        celda.font = encabezado_font
        celda.alignment = Alignment(vertical="center")

    for fila_idx, fila in enumerate(filas, start=2):
        for col_idx, valor in enumerate(fila, start=1):
            celda = hoja.cell(row=fila_idx, column=col_idx, value=valor)
            if isinstance(valor, float):
                celda.number_format = "#,##0.00"

    for col_idx, titulo in enumerate(encabezados, start=1):
        largo = len(str(titulo))
        for fila in filas[:300]:
            if col_idx - 1 < len(fila):
                largo = max(largo, len(str(fila[col_idx - 1])))
        hoja.column_dimensions[hoja.cell(row=1, column=col_idx).column_letter].width = min(max(largo + 2, 10), 40)
    hoja.freeze_panes = "A2"
    if filas:
        hoja.auto_filter.ref = f"A1:{hoja.cell(row=1, column=len(encabezados)).column_letter}{len(filas) + 1}"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )


def _datos_reporte(tipo, fecha_desde, fecha_hasta, tecnico_id, estado, db):
    """Arma las filas y encabezados de un reporte, en el mismo formato
    para CSV y para Excel. Devuelve (filas, encabezados, nombre_base,
    titulo_hoja)."""
    if tipo in ("ordenes", "reparaciones"):
        ordenes = _filtrar_ordenes(db, fecha_desde, fecha_hasta, tecnico_id, estado)
        filas = [[o.numero_orden, o.fecha.strftime("%d/%m/%Y"), o.cliente.nombre if o.cliente else "",
                  o.tipo_equipo, o.marca, o.modelo, o.estado, o.tecnico.nombre if o.tecnico else "",
                  float(o.total), float(o.abonado), float(o.saldo)] for o in ordenes]
        return filas, ["N. Orden", "Fecha", "Cliente", "Tipo", "Marca", "Modelo", "Estado", "Tecnico", "Total", "Abonado", "Saldo"], "reporte_ordenes", "Ordenes"

    if tipo in ("ingresos", "gastos"):
        q = db.query(MovimientoFinanciero).filter(MovimientoFinanciero.tipo == ("ingreso" if tipo == "ingresos" else "gasto"))
        if fecha_desde:
            q = q.filter(MovimientoFinanciero.fecha >= fecha_desde)
        if fecha_hasta:
            q = q.filter(MovimientoFinanciero.fecha <= fecha_hasta)
        movimientos = q.order_by(MovimientoFinanciero.fecha.desc()).all()
        filas = [[m.fecha.strftime("%d/%m/%Y %H:%M"), m.categoria, float(m.monto), m.descripcion, m.referencia] for m in movimientos]
        return filas, ["Fecha", "Categoria", "Monto", "Descripcion", "Referencia"], f"reporte_{tipo}", ("Ingresos" if tipo == "ingresos" else "Gastos")

    if tipo == "ventas":
        ventas = db.query(Venta).order_by(Venta.fecha.desc()).all()
        filas = [[v.numero_venta, v.fecha.strftime("%d/%m/%Y %H:%M"), float(v.subtotal), float(v.descuento),
                  float(v.total), v.forma_pago] for v in ventas]
        return filas, ["N. Venta", "Fecha", "Subtotal", "Descuento", "Total", "Forma de pago"], "reporte_ventas", "Ventas"

    if tipo == "inventario":
        productos = db.query(Producto).order_by(Producto.nombre).all()
        filas = [[p.codigo, p.nombre, p.categoria, p.marca, p.caja, float(p.costo), float(p.precio_venta),
                  p.existencia, p.stock_minimo] for p in productos]
        return filas, ["Codigo", "Nombre", "Categoria", "Marca", "Caja", "Costo", "Precio", "Existencia", "Stock minimo"], "reporte_inventario", "Inventario"

    if tipo == "saldos":
        ordenes = db.query(OrdenServicio).options(joinedload(OrdenServicio.cliente)).filter(
            OrdenServicio.saldo > 0, OrdenServicio.estado != "CANCELADO"
        ).order_by(OrdenServicio.fecha.desc()).all()
        filas = [[o.numero_orden, o.cliente.nombre if o.cliente else "", float(o.total), float(o.abonado), float(o.saldo)] for o in ordenes]
        return filas, ["N. Orden", "Cliente", "Total", "Abonado", "Saldo"], "reporte_saldos", "Saldos"

    return [], ["Sin datos"], "reporte", "Reporte"


def _rango_fechas(desde, hasta):
    fecha_desde = _parse_fecha(desde)
    fecha_hasta = _parse_fecha(hasta)
    if fecha_hasta:
        fecha_hasta = fecha_hasta.replace(hour=23, minute=59, second=59)
    return fecha_desde, fecha_hasta


@router.get("/reportes/exportar/{tipo}")
def reportes_exportar_csv(
    tipo: str, desde: str = "", hasta: str = "", tecnico_id: str = "", estado: str = "", forma_pago: str = "",
    db: Session = Depends(get_db), usuario=Depends(login_required),
):
    fecha_desde, fecha_hasta = _rango_fechas(desde, hasta)
    filas, encabezados, nombre_base, _ = _datos_reporte(tipo, fecha_desde, fecha_hasta, tecnico_id, estado, db)
    return _csv_response(filas, encabezados, f"{nombre_base}.csv")


@router.get("/reportes/exportar/{tipo}/excel")
def reportes_exportar_excel(
    tipo: str, desde: str = "", hasta: str = "", tecnico_id: str = "", estado: str = "", forma_pago: str = "",
    db: Session = Depends(get_db), usuario=Depends(login_required),
):
    fecha_desde, fecha_hasta = _rango_fechas(desde, hasta)
    filas, encabezados, nombre_base, titulo_hoja = _datos_reporte(tipo, fecha_desde, fecha_hasta, tecnico_id, estado, db)
    return _xlsx_response(filas, encabezados, f"{nombre_base}.xlsx", titulo_hoja)
