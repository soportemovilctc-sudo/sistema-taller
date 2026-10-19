"""Mejoras tomadas de Zafra: Bitácora, Proveedores, Cotizaciones, Cuentas
por cobrar con antigüedad, Declaración de ISV y Reorden de inventario.
Los cálculos están en app/utils/mejoras.py."""
from datetime import datetime, timedelta
from io import BytesIO
from urllib.parse import urlencode

from fastapi import APIRouter, Request, Depends, Form
from fastapi.responses import RedirectResponse, Response
from sqlalchemy import or_, func
from sqlalchemy.orm import Session

from app.templates_env import templates, obtener_config_actual
from app.database import get_db
from app.deps import roles_required, login_required
from app.models import (
    Bitacora, Proveedor, Cotizacion, Cliente, Producto, SolicitudCompra, RecepcionCompra, Categoria,
    COTIZACION_VIGENTE, COTIZACION_ANULADA, ESTADOS_COTIZACION,
)
from app.utils.flash import flash
from app.utils.calculations import to_decimal
from app.utils.turnos_caja import hoy_local, inicio_dia_local_en_utc, a_hora_local
from app.utils import mejoras as M

router = APIRouter()
solo_admin = roles_required("admin")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MESES = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def _fecha(t):
    try:
        return datetime.strptime((t or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _excel(titulo: str, encabezados: list, filas: list, nombre: str, resumen: list | None = None) -> Response:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    h = wb.active
    h.title = titulo[:30]
    cfg = obtener_config_actual()
    h.append([cfg.get("nombre_taller", "")])
    h["A1"].font = Font(bold=True, size=13)
    h.append([titulo])
    h["A2"].font = Font(bold=True)
    for r in (resumen or []):
        h.append(r)
    h.append([])
    h.append(encabezados)
    for c in h[h.max_row]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    for f in filas:
        h.append([float(v) if hasattr(v, "quantize") else v for v in f])
    for col in "ABCDEFGHIJKL":
        h.column_dimensions[col].width = 18
    h.column_dimensions["B"].width = 34
    out = BytesIO()
    wb.save(out)
    return Response(out.getvalue(), media_type=XLSX, headers={"Content-Disposition": f'attachment; filename="{nombre}.xlsx"'})


# ---------------------------------------------------------------------------
# Bitácora
# ---------------------------------------------------------------------------
@router.get("/bitacora")
def bitacora(request: Request, desde: str = "", hasta: str = "", usuario_f: str = "", modulo: str = "", q: str = "",
             db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    d, h = _fecha(desde), _fecha(hasta)
    if not d and not h:
        h = hoy_local()
        d = h - timedelta(days=6)
    consulta = db.query(Bitacora)
    if d:
        consulta = consulta.filter(Bitacora.fecha >= inicio_dia_local_en_utc(d))
    if h:
        consulta = consulta.filter(Bitacora.fecha < inicio_dia_local_en_utc(h + timedelta(days=1)))
    if usuario_f:
        consulta = consulta.filter(Bitacora.usuario_nombre == usuario_f)
    if modulo:
        consulta = consulta.filter(Bitacora.modulo == modulo)
    if q.strip():
        consulta = consulta.filter(Bitacora.accion.ilike(f"%{q.strip()}%"))
    total = consulta.count()
    filas = consulta.order_by(Bitacora.fecha.desc()).limit(500).all()
    usuarios = [u for (u,) in db.query(Bitacora.usuario_nombre).distinct().order_by(Bitacora.usuario_nombre).all() if u]
    modulos = [m for (m,) in db.query(Bitacora.modulo).distinct().order_by(Bitacora.modulo).all() if m]
    return templates.TemplateResponse("mejoras/bitacora.html", {
        "request": request, "usuario": usuario, "filas": filas, "total": total, "usuarios": usuarios, "modulos": modulos,
        "desde": d.isoformat() if d else "", "hasta": h.isoformat() if h else "", "usuario_f": usuario_f, "modulo": modulo, "q": q,
        "hoy_iso": hoy_local().isoformat(),
    })


# ---------------------------------------------------------------------------
# Proveedores
# ---------------------------------------------------------------------------
def _saldos_proveedor(db: Session) -> dict:
    """nombre en minúsculas -> {comprado, saldo, recepciones, ultima}"""
    out = {}
    for r in db.query(RecepcionCompra).filter(RecepcionCompra.anulada == False).all():  # noqa: E712
        k = (r.proveedor or "").strip().lower()
        x = out.setdefault(k, {"comprado": M.CERO, "saldo": M.CERO, "recepciones": 0, "ultima": None})
        x["comprado"] += M.q(r.total)
        x["saldo"] += M.q(r.saldo)
        x["recepciones"] += 1
        x["ultima"] = max(x["ultima"], r.fecha) if x["ultima"] else r.fecha
    return out


def _datos_proveedor(form) -> dict:
    return {k: (form.get(k) or "").strip() for k in ("nombre", "rtn", "telefono", "whatsapp", "correo", "direccion", "contacto", "notas")}


@router.get("/proveedores")
def proveedores(request: Request, q: str = "", db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    consulta = db.query(Proveedor)
    if q.strip():
        like = f"%{q.strip()}%"
        consulta = consulta.filter(or_(Proveedor.nombre.ilike(like), Proveedor.rtn.ilike(like), Proveedor.telefono.ilike(like)))
    lista = consulta.order_by(Proveedor.activo.desc(), Proveedor.nombre).all()
    saldos = _saldos_proveedor(db)
    registrados = {p.nombre.strip().lower() for p in db.query(Proveedor).all()}
    sin_registrar = sorted({(r.proveedor or "").strip() for r in db.query(SolicitudCompra).all()
                            if (r.proveedor or "").strip() and (r.proveedor or "").strip().lower() not in registrados})
    return templates.TemplateResponse("mejoras/proveedores.html", {
        "request": request, "usuario": usuario, "proveedores": lista, "saldos": saldos, "q": q, "sin_registrar": sin_registrar,
        "total_saldo": sum((v["saldo"] for k, v in saldos.items()), M.CERO),
    })


@router.post("/proveedores/nuevo")
async def proveedores_nuevo(request: Request, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    datos = _datos_proveedor(await request.form())
    if not datos["nombre"]:
        flash(request, "Escribe el nombre del proveedor.", "error")
        return RedirectResponse("/proveedores", status_code=303)
    if db.query(Proveedor).filter(func.lower(Proveedor.nombre) == datos["nombre"].lower()).first():
        flash(request, "Ya existe un proveedor con ese nombre.", "error")
        return RedirectResponse("/proveedores", status_code=303)
    p = Proveedor(**datos, activo=True)
    db.add(p)
    db.flush()
    db.query(SolicitudCompra).filter(func.lower(SolicitudCompra.proveedor) == datos["nombre"].lower(),
                                     SolicitudCompra.proveedor_id.is_(None)).update({"proveedor_id": p.id}, synchronize_session=False)
    db.commit()
    flash(request, f"Proveedor «{p.nombre}» guardado.", "success")
    return RedirectResponse(f"/proveedores/{p.id}", status_code=303)


@router.get("/proveedores/{proveedor_id}")
def proveedor_detalle(proveedor_id: int, request: Request, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    p = db.get(Proveedor, proveedor_id)
    if not p:
        return RedirectResponse("/proveedores", status_code=303)
    solicitudes = (db.query(SolicitudCompra).filter(or_(SolicitudCompra.proveedor_id == p.id, func.lower(SolicitudCompra.proveedor) == p.nombre.lower()))
                   .order_by(SolicitudCompra.id.desc()).limit(100).all())
    recepciones = [r for s in solicitudes for r in s.recepciones if not r.anulada]
    return templates.TemplateResponse("mejoras/proveedor_detalle.html", {
        "request": request, "usuario": usuario, "p": p, "solicitudes": solicitudes,
        "comprado": sum((M.q(r.total) for r in recepciones), M.CERO), "saldo": sum((M.q(r.saldo) for r in recepciones), M.CERO),
        "n_recepciones": len(recepciones),
    })


@router.post("/proveedores/{proveedor_id}/editar")
async def proveedor_editar(proveedor_id: int, request: Request, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    p = db.get(Proveedor, proveedor_id)
    if not p:
        return RedirectResponse("/proveedores", status_code=303)
    form = await request.form()
    datos = _datos_proveedor(form)
    if not datos["nombre"]:
        flash(request, "El nombre no puede quedar vacío.", "error")
        return RedirectResponse(f"/proveedores/{p.id}", status_code=303)
    for k, v in datos.items():
        setattr(p, k, v)
    p.activo = form.get("activo") != "no"
    db.commit()
    flash(request, "Proveedor actualizado.", "success")
    return RedirectResponse(f"/proveedores/{p.id}", status_code=303)


# ---------------------------------------------------------------------------
# Cotizaciones
# ---------------------------------------------------------------------------
def _tasa(db):
    from app.routers.inventario import _isv_tasa
    return _isv_tasa(db)


@router.get("/cotizaciones")
def cotizaciones(request: Request, estado: str = "", q: str = "", db: Session = Depends(get_db), usuario=Depends(login_required)):
    consulta = db.query(Cotizacion)
    if q.strip():
        like = f"%{q.strip()}%"
        consulta = consulta.filter(or_(Cotizacion.numero.ilike(like), Cotizacion.cliente_nombre.ilike(like)))
    lista = consulta.order_by(Cotizacion.id.desc()).limit(300).all()
    if estado:
        lista = [c for c in lista if c.estado_actual == estado]
    vigentes = [c for c in lista if c.estado_actual == COTIZACION_VIGENTE]
    return templates.TemplateResponse("mejoras/cotizaciones.html", {
        "request": request, "usuario": usuario, "cotizaciones": lista, "estado": estado, "q": q, "estados": ESTADOS_COTIZACION,
        "n_vigentes": len(vigentes), "total_vigentes": sum((M.q(c.total) for c in vigentes), M.CERO),
    })


@router.get("/cotizaciones/nueva")
def cotizacion_nueva(request: Request, db: Session = Depends(get_db), usuario=Depends(login_required)):
    productos = db.query(Producto).filter(Producto.estado == "activo").order_by(Producto.nombre).all()
    from app.models import Servicio
    servicios = db.query(Servicio).order_by(Servicio.nombre).all()
    return templates.TemplateResponse("mejoras/cotizacion_nueva.html", {
        "request": request, "usuario": usuario, "productos": productos, "servicios": servicios,
        "clientes": db.query(Cliente).order_by(Cliente.nombre).all(), "tasa_isv": _tasa(db),
    })


@router.post("/cotizaciones/nueva")
async def cotizacion_crear(request: Request, db: Session = Depends(get_db), usuario=Depends(login_required)):
    form = await request.form()
    ids, descs, cants, precios = form.getlist("producto_id"), form.getlist("descripcion"), form.getlist("cantidad"), form.getlist("precio")
    lineas = [{"producto_id": ids[i] if i < len(ids) else "", "descripcion": descs[i] if i < len(descs) else "",
               "cantidad": cants[i] if i < len(cants) else "1", "precio": precios[i] if i < len(precios) else "0"} for i in range(len(descs))]
    datos = {k: form.get(k, "") for k in ("cliente_id", "cliente_nombre", "cliente_telefono", "validez_dias", "notas")}
    datos["precios_con_isv"] = bool(form.get("precios_con_isv"))
    cot, errores = M.crear_cotizacion(db, datos, lineas, usuario["nombre_completo"], _tasa(db))
    if errores:
        db.rollback()
        flash(request, "No se guardó la cotización: " + " ".join(errores), "error")
        return RedirectResponse("/cotizaciones/nueva", status_code=303)
    db.commit()
    flash(request, f"Cotización {cot.numero} guardada. Ya puedes mandarla por WhatsApp o en PDF.", "success")
    return RedirectResponse(f"/cotizaciones/{cot.id}", status_code=303)


@router.get("/cotizaciones/{cot_id}")
def cotizacion_detalle(cot_id: int, request: Request, db: Session = Depends(get_db), usuario=Depends(login_required)):
    cot = db.get(Cotizacion, cot_id)
    if not cot:
        return RedirectResponse("/cotizaciones", status_code=303)
    cfg = obtener_config_actual()
    tel = cot.cliente_telefono or ((cot.cliente.whatsapp or cot.cliente.telefono) if cot.cliente else "")
    tasa = (to_decimal(cot.isv) / to_decimal(cot.subtotal) * 100) if to_decimal(cot.subtotal or 0) > 0 else to_decimal(0)
    lineas = [{"d": d, "unit_con": M.q(to_decimal(d.precio_unitario) * (1 + tasa / 100)),
               "total_con": M.q(to_decimal(d.precio_unitario) * (1 + tasa / 100) * d.cantidad)} for d in cot.detalles]
    from app.models import Venta
    return templates.TemplateResponse("mejoras/cotizacion_detalle.html", {
        "request": request, "usuario": usuario, "cot": cot, "lineas": lineas, "estados": ESTADOS_COTIZACION,
        "whatsapp": M.link_whatsapp(tel, M.mensaje_whatsapp_cotizacion(cot, cfg.get("nombre_taller", ""), cfg.get("moneda", "L"))),
        "telefono": tel, "venta": db.get(Venta, cot.venta_id) if cot.venta_id else None,
        "todo_productos": all(d.producto_id for d in cot.detalles),
    })


@router.get("/cotizaciones/{cot_id}/pdf")
def cotizacion_pdf(cot_id: int, descargar: int = 0, db: Session = Depends(get_db), usuario=Depends(login_required)):
    from app.models import Configuracion, ConfiguracionFacturacion
    from app.utils.pdf_solicitud_compra import pdf_cotizacion
    cot = db.get(Cotizacion, cot_id)
    if not cot:
        return RedirectResponse("/cotizaciones", status_code=303)
    contenido = pdf_cotizacion(cot, db.get(Configuracion, 1), db.query(ConfiguracionFacturacion).first())
    modo = "attachment" if descargar else "inline"
    return Response(contenido, media_type="application/pdf", headers={"Content-Disposition": f'{modo}; filename="cotizacion_{cot.numero}.pdf"'})


@router.post("/cotizaciones/{cot_id}/anular")
def cotizacion_anular(cot_id: int, request: Request, db: Session = Depends(get_db), usuario=Depends(login_required)):
    cot = db.get(Cotizacion, cot_id)
    if cot and cot.estado == COTIZACION_VIGENTE:
        cot.estado = COTIZACION_ANULADA
        db.commit()
        flash(request, f"Cotización {cot.numero} anulada.", "success")
    return RedirectResponse(f"/cotizaciones/{cot_id}", status_code=303)


# ---------------------------------------------------------------------------
# Cuentas por cobrar con antigüedad
# ---------------------------------------------------------------------------
@router.get("/cuentas-por-cobrar")
def cxc(request: Request, tramo: str = "", q: str = "", db: Session = Depends(get_db), usuario=Depends(roles_required("vendedor"))):
    datos = M.cuentas_por_cobrar(db, tramo, q)
    return templates.TemplateResponse("mejoras/cxc.html", {"request": request, "usuario": usuario, "d": datos, "tramo": tramo, "q": q})


@router.get("/cuentas-por-cobrar/excel")
def cxc_excel(tramo: str = "", q: str = "", db: Session = Depends(get_db), usuario=Depends(roles_required("vendedor"))):
    d = M.cuentas_por_cobrar(db, tramo, q)
    filas = [[f["orden"].numero_orden, f["cliente"], f["telefono"], a_hora_local(f["orden"].fecha).strftime("%d/%m/%Y"), f["dias"],
              f["tramo"], f["orden"].estado, f["total"], f["abonado"], f["saldo"]] for f in d["filas"]]
    resumen = [[v["nombre"], v["n"], float(v["saldo"])] for v in d["tramos"].values()] + [["TOTAL", d["n"], float(d["total"])]]
    return _excel("Cuentas por cobrar", ["Orden", "Cliente", "Teléfono", "Fecha", "Días", "Antigüedad", "Estado", "Total", "Abonado", "Saldo"],
                  filas, "cuentas_por_cobrar", resumen)


# ---------------------------------------------------------------------------
# Declaración de ISV
# ---------------------------------------------------------------------------
def _mes(anio, mes):
    hoy = hoy_local()
    try:
        anio, mes = int(anio or hoy.year), int(mes or hoy.month)
        if not 1 <= mes <= 12:
            raise ValueError
    except ValueError:
        anio, mes = hoy.year, hoy.month
    return anio, mes


@router.get("/impuestos/isv")
def isv(request: Request, anio: str = "", mes: str = "", db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    anio, mes = _mes(anio, mes)
    hoy = hoy_local()
    return templates.TemplateResponse("mejoras/isv.html", {
        "request": request, "usuario": usuario, "d": M.declaracion_isv(db, anio, mes), "meses": MESES,
        "anios": list(range(hoy.year - 3, hoy.year + 1)), "a_hora_local": a_hora_local,
    })


@router.get("/impuestos/isv/excel")
def isv_excel(anio: str = "", mes: str = "", db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    anio, mes = _mes(anio, mes)
    d = M.declaracion_isv(db, anio, mes)
    filas = [["Venta", f.numero_documento, a_hora_local(f.fecha_emision).strftime("%d/%m/%Y"), f.cliente_nombre, f.cliente_rtn if (f.cliente_rtn or "").lower() != "none" else "",
              "Fiscal" if f.tipo == "fiscal" else "Interna", M.q(f.importe_exento), M.q(f.importe_gravado), M.q(f.isv_monto), M.q(f.total)]
             for f in d["lista_fiscales"] + d["lista_internas"]]
    filas += [["Compra", r.factura_proveedor or r.numero, a_hora_local(r.fecha).strftime("%d/%m/%Y"), r.proveedor, "", "Crédito fiscal",
               0, M.q(r.subtotal), M.q(r.isv), M.q(r.total)] for r in d["lista_compras"]]
    resumen = [["ISV de ventas (facturas fiscales)", float(d["fiscales"]["isv"])], ["ISV de compras (crédito fiscal)", float(d["compras"]["isv"])],
               ["ISV a pagar", float(d["a_pagar"])], ["ISV de comprobantes internos (aparte)", float(d["internas"]["isv"])]]
    return _excel(f"ISV {MESES[mes]} {anio}", ["Tipo", "Documento", "Fecha", "Cliente / Proveedor", "RTN", "Clase", "Exento", "Gravado / Base", "ISV", "Total"],
                  filas, f"declaracion_isv_{anio}_{mes:02d}", resumen)


# ---------------------------------------------------------------------------
# Reorden de inventario
# ---------------------------------------------------------------------------
@router.get("/inventario/reorden")
def reorden(request: Request, categoria: str = "", db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    filas = M.reorden(db, categoria)
    cats = [c.nombre for c in db.query(Categoria).filter(Categoria.activo == True).order_by(Categoria.nombre).all()]  # noqa: E712
    sin_minimo = db.query(Producto).filter(Producto.estado == "activo", (Producto.stock_minimo == None) | (Producto.stock_minimo <= 0)).count()  # noqa: E711
    return templates.TemplateResponse("mejoras/reorden.html", {
        "request": request, "usuario": usuario, "filas": filas, "categoria": categoria, "categorias": cats,
        "proveedores": M.proveedores_activos(db), "sin_minimo": sin_minimo,
        "total_valor": sum((f["valor"] for f in filas), M.CERO),
    })


@router.post("/inventario/reorden/solicitud")
async def reorden_solicitud(request: Request, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    from app.utils.compras import crear_solicitud
    form = await request.form()
    elegidos = form.getlist("elegir")
    lineas = []
    for pid in elegidos:
        cant = (form.get(f"cantidad_{pid}") or "0").strip()
        p = db.get(Producto, int(pid)) if pid.isdigit() else None
        if p and cant.isdigit() and int(cant) > 0:
            lineas.append({"producto_id": str(p.id), "cantidad": cant, "costo": f"{to_decimal(p.costo or 0):.2f}", "precio": ""})
    proveedor = (form.get("proveedor") or "").strip()
    if not lineas:
        flash(request, "Marca al menos un producto con cantidad mayor a 0.", "error")
        return RedirectResponse("/inventario/reorden", status_code=303)
    sol, errores = crear_solicitud(db, proveedor, "Creada desde el reorden de inventario", lineas, usuario["nombre_completo"])
    if errores:
        db.rollback()
        flash(request, "No se creó la solicitud: " + " ".join(errores), "error")
        return RedirectResponse("/inventario/reorden", status_code=303)
    prov = db.query(Proveedor).filter(func.lower(Proveedor.nombre) == proveedor.lower()).first() if proveedor else None
    if prov:
        sol.proveedor_id = prov.id
    db.commit()
    flash(request, f"Solicitud {sol.numero} creada con {len(lineas)} producto(s). Revísala y descárgala en PDF para el proveedor.", "success")
    return RedirectResponse(f"/compras/{sol.id}", status_code=303)
