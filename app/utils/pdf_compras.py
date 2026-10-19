"""Documentos de Compras: el comprobante de una recepción de mercadería
(RC-000001) y el reporte de recepciones de un período, en PDF y Excel.
Respaldan cada ingreso de producto al inventario con cantidades, costos y
la existencia antes y después."""
from datetime import datetime
from decimal import Decimal
from io import BytesIO
from xml.sax.saxutils import escape

from sqlalchemy.orm import Session, selectinload

from app.models import RecepcionCompra, RecepcionCompraDetalle, MovimientoFinanciero
from app.utils.calculations import to_decimal
from app.utils.turnos_caja import a_hora_local
from app.utils.pdf_ajustes import _estilos, _tabla, _resumen, _m

CENT = Decimal("0.01")
CERO = Decimal("0.00")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _f(dt, formato="%d/%m/%Y %I:%M %p"):
    return a_hora_local(dt).strftime(formato) if dt else ""


def _d(v) -> Decimal:
    return to_decimal(v or 0).quantize(CENT)


def _producto_txt(p) -> str:
    if not p:
        return "(producto eliminado)"
    return p.nombre + (f" ({p.categoria})" if p.categoria else "")


def pagos_de_recepcion(db: Session, recepcion_id: int) -> list:
    return (db.query(MovimientoFinanciero).filter(MovimientoFinanciero.recepcion_compra_id == recepcion_id)
            .order_by(MovimientoFinanciero.fecha).all())


def _forma_txt(forma: str) -> str:
    return {"Crédito": "Al crédito (por pagar al proveedor)", "Efectivo": "Efectivo de caja", "Banco": "Banco (transferencia)"}.get(forma, forma)


def _origen_pago(m) -> str:
    return "Efectivo de caja" if m.cuenta == "caja_chica" else "Banco"


def _lineas(rec):
    """Filas del comprobante: cantidad, costo, subtotal y existencia antes/después."""
    filas = []
    for d in rec.detalles:
        antes = d.existencia_anterior or 0
        filas.append({"codigo": d.producto.codigo if d.producto else "", "producto": _producto_txt(d.producto),
                      "cantidad": d.cantidad, "costo": _d(d.costo_unitario), "subtotal": _d(d.subtotal),
                      "antes": antes, "despues": antes + (d.cantidad or 0)})
    return filas


# ---------------------------------------------------------------------------
# Comprobante de una recepción
# ---------------------------------------------------------------------------
def pdf_recepcion(rec, pagos: list, nombre_taller: str, moneda: str) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    e = _estilos()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, topMargin=1.3 * cm, bottomMargin=1.3 * cm,
                            leftMargin=1.4 * cm, rightMargin=1.4 * cm, title=f"Recepción {rec.numero}")
    filas = _lineas(rec)
    unidades = sum(f["cantidad"] for f in filas)
    sol = rec.solicitud
    story = [Paragraph(escape(nombre_taller), e["titulo"]),
             Paragraph(f"Comprobante de ingreso de mercadería <b>{escape(rec.numero)}</b> · {_f(rec.fecha)}", e["sub"]),
             Paragraph(f"Proveedor: <b>{escape(rec.proveedor or '—')}</b>"
                       + (f" · Factura del proveedor: <b>{escape(rec.factura_proveedor)}</b>" if rec.factura_proveedor else " · Sin número de factura del proveedor")
                       + (f" · Solicitud {escape(sol.numero)}" if sol else ""), e["sub"]),
             Paragraph(f"Recibido por {escape(rec.usuario_nombre or '')} · Pago: {escape(_forma_txt(rec.forma_pago))}", e["sub"])]
    if rec.notas:
        story.append(Paragraph(f"Notas: {escape(rec.notas)}", e["sub"]))
    if getattr(rec, "anulada", False):
        story.append(Paragraph(
            f"<font color='#c0392b'><b>ANULADA</b> el {_f(rec.anulada_en)} por {escape(rec.anulada_por or '')}"
            + (f" · Motivo: {escape(rec.motivo_anulacion)}" if rec.motivo_anulacion else "")
            + ". Las unidades se sacaron del inventario y los pagos se quitaron de la contabilidad.</font>", e["sub"]))
    story += [Spacer(1, 8),
              _resumen([("Unidades que entraron", f"{unidades}"), ("Mercadería (sin ISV)", _m(rec.subtotal, moneda)),
                        ("ISV de la factura", _m(rec.isv, moneda)), ("Total de la compra", _m(rec.total, moneda))], moneda),
              Paragraph("Productos recibidos", e["seccion"])]
    datos_t = [["Código", "Producto", "Cant.", "Costo c/u", "Subtotal", "Había", "Quedó"]]
    for f in filas:
        datos_t.append([Paragraph(escape(f["codigo"]), e["celda"]), Paragraph(escape(f["producto"]), e["celda"]), f["cantidad"],
                        _m(f["costo"], moneda), _m(f["subtotal"], moneda), f["antes"], f["despues"]])
    datos_t.append(["", "Mercadería (sin ISV)", unidades, "", _m(rec.subtotal, moneda), "", ""])
    if rec.isv:
        datos_t.append(["", "ISV de la factura", "", "", _m(rec.isv, moneda), "", ""])
    datos_t.append(["", "Total", "", "", _m(rec.total, moneda), "", ""])
    t = _tabla(datos_t, [2.6 * cm, 6.4 * cm, 1.4 * cm, 2.4 * cm, 2.6 * cm, 1.4 * cm, 1.4 * cm], num_cols=(2, 3, 4, 5, 6), negrita_ultima=True)
    story.append(t)
    story.append(Spacer(1, 3))
    story.append(Paragraph("Costos sin ISV. «Había» y «Quedó» son las existencias del producto justo antes y después de esta recepción.", e["sub"]))

    story.append(Paragraph("Pagos al proveedor", e["seccion"]))
    datos_p = [["Fecha", "De dónde salió", "Registrado por", "Monto"]]
    for m in pagos:
        datos_p.append([_f(m.fecha), _origen_pago(m), escape(m.usuario_nombre or ""),
                        _m(m.monto, moneda) + (" (anulado)" if getattr(m, "excluir_de_contabilidad", False) else "")])
    if len(datos_p) == 1:
        datos_p.append(["—", "Sin pagos todavía", "", ""])
    datos_p.append(["", "", "Pagado (anulado)" if getattr(rec, "anulada", False) else "Pagado", _m(rec.pagado, moneda)])
    datos_p.append(["", "", "Saldo pendiente", _m(rec.saldo, moneda)])
    story.append(_tabla(datos_p, [4.2 * cm, 4.6 * cm, 5.0 * cm, 3.0 * cm], num_cols=(3,), negrita_ultima=True))
    story += [Spacer(1, 14), Paragraph(f"Documento generado el {_f(datetime.utcnow())}.", e["sub"])]
    doc.build(story)
    return buffer.getvalue()


def excel_recepcion(rec, pagos: list, nombre_taller: str, moneda: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    h = wb.active
    h.title = rec.numero
    sol = rec.solicitud
    h["A1"] = nombre_taller
    h["A1"].font = Font(size=14, bold=True)
    h["A2"] = f"Comprobante de ingreso de mercadería {rec.numero} · {_f(rec.fecha)}"
    h["A2"].font = Font(size=12, bold=True)
    h["A3"] = f"Proveedor: {rec.proveedor or '—'} · Factura del proveedor: {rec.factura_proveedor or '—'}" + (f" · Solicitud {sol.numero}" if sol else "")
    h["A4"] = f"Recibido por {rec.usuario_nombre or ''} · Pago: {_forma_txt(rec.forma_pago)}"
    if rec.notas:
        h["A5"] = f"Notas: {rec.notas}"
    fila = 7
    for col, t in enumerate(["Código", "Producto", "Cantidad", f"Costo c/u ({moneda})", f"Subtotal ({moneda})", "Había", "Quedó"], start=1):
        c = h.cell(row=fila, column=col, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
    filas = _lineas(rec)
    for f in filas:
        fila += 1
        for col, v in enumerate([f["codigo"], f["producto"], f["cantidad"], float(f["costo"]), float(f["subtotal"]), f["antes"], f["despues"]], start=1):
            c = h.cell(row=fila, column=col, value=v)
            if col in (4, 5):
                c.number_format = '#,##0.00'
    for etiqueta, valor, unidades in (("Mercadería (sin ISV)", rec.subtotal, sum(f["cantidad"] for f in filas)),
                                      ("ISV de la factura", rec.isv, None), ("Total", rec.total, None)):
        fila += 1
        h.cell(row=fila, column=2, value=etiqueta).font = Font(bold=True)
        if unidades is not None:
            h.cell(row=fila, column=3, value=unidades).font = Font(bold=True)
        c = h.cell(row=fila, column=5, value=float(_d(valor)))
        c.font = Font(bold=True)
        c.number_format = '#,##0.00'
    fila += 2
    h.cell(row=fila, column=1, value="Pagos al proveedor").font = Font(bold=True)
    for m in pagos:
        fila += 1
        h.cell(row=fila, column=1, value=_f(m.fecha))
        h.cell(row=fila, column=2, value=f"{_origen_pago(m)} · {m.usuario_nombre or ''}")
        c = h.cell(row=fila, column=5, value=float(_d(m.monto)))
        c.number_format = '#,##0.00'
    for etiqueta, valor in (("Pagado", rec.pagado), ("Saldo pendiente", rec.saldo)):
        fila += 1
        h.cell(row=fila, column=2, value=etiqueta).font = Font(bold=True)
        c = h.cell(row=fila, column=5, value=float(_d(valor)))
        c.font = Font(bold=True)
        c.number_format = '#,##0.00'
    for col, ancho in zip("ABCDEFG", (22, 46, 11, 15, 15, 9, 9)):
        h.column_dimensions[col].width = ancho
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# Reporte de recepciones de un período
# ---------------------------------------------------------------------------
def _acum():
    return {"recepciones": 0, "unidades": 0, "subtotal": CERO, "isv": CERO, "total": CERO, "pagado": CERO, "saldo": CERO}


def _sumar(a, rec, unidades):
    a["recepciones"] += 1
    a["unidades"] += unidades
    a["subtotal"] += _d(rec.subtotal)
    a["isv"] += _d(rec.isv)
    a["total"] += _d(rec.total)
    a["pagado"] += _d(rec.pagado)
    a["saldo"] += _d(rec.saldo)


def reporte_recepciones(db: Session, inicio, fin, proveedor: str = "") -> dict:
    """Recepciones con fecha en [inicio, fin) (UTC), opcionalmente de un proveedor."""
    consulta = db.query(RecepcionCompra).options(
        selectinload(RecepcionCompra.detalles).selectinload(RecepcionCompraDetalle.producto),
        selectinload(RecepcionCompra.solicitud))
    if inicio:
        consulta = consulta.filter(RecepcionCompra.fecha >= inicio)
    if fin:
        consulta = consulta.filter(RecepcionCompra.fecha < fin)
    if proveedor.strip():
        consulta = consulta.filter(RecepcionCompra.proveedor.ilike(f"%{proveedor.strip()}%"))
    consulta = consulta.filter(RecepcionCompra.anulada == False)  # noqa: E712
    recepciones = consulta.order_by(RecepcionCompra.fecha, RecepcionCompra.id).all()

    total, por_proveedor, por_forma, por_producto, filas, lista = _acum(), {}, {}, {}, [], []
    for r in recepciones:
        unidades = sum(d.cantidad or 0 for d in r.detalles)
        _sumar(total, r, unidades)
        _sumar(por_proveedor.setdefault(r.proveedor or "(sin proveedor)", _acum()), r, unidades)
        _sumar(por_forma.setdefault(r.forma_pago, _acum()), r, unidades)
        lista.append({"rec": r, "unidades": unidades})
        for d in r.detalles:
            antes = d.existencia_anterior or 0
            p = d.producto
            filas.append({"fecha": r.fecha, "rec": r, "codigo": p.codigo if p else "", "nombre": p.nombre if p else "(eliminado)",
                          "categoria": (p.categoria or "") if p else "", "cantidad": d.cantidad, "costo": _d(d.costo_unitario),
                          "subtotal": _d(d.subtotal), "antes": antes, "despues": antes + (d.cantidad or 0)})
            g = por_producto.setdefault(d.producto_id, {"codigo": p.codigo if p else "", "nombre": p.nombre if p else "(eliminado)",
                                                       "categoria": (p.categoria or "") if p else "", "unidades": 0, "valor": CERO})
            g["unidades"] += d.cantidad or 0
            g["valor"] += _d(d.subtotal)
    for g in por_producto.values():
        g["costo_promedio"] = (g["valor"] / g["unidades"]).quantize(CENT) if g["unidades"] else CERO

    def ordenar(dic):
        return [dict(v, clave=k) for k, v in sorted(dic.items(), key=lambda kv: -kv[1]["total"])]

    return {"total": total, "recepciones": lista, "filas": filas,
            "por_proveedor": ordenar(por_proveedor), "por_forma": ordenar(por_forma),
            "por_producto": sorted(por_producto.values(), key=lambda g: (-g["valor"], g["nombre"]))}


def pdf_reporte_recepciones(datos: dict, rango_txt: str, nombre_taller: str, moneda: str, filtros: str = "") -> bytes:
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    e = _estilos()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=1.1 * cm, bottomMargin=1.1 * cm,
                            leftMargin=1.2 * cm, rightMargin=1.2 * cm, title="Reporte de recepciones de compra")
    t = datos["total"]
    story = [Paragraph(escape(nombre_taller), e["titulo"]),
             Paragraph(f"Reporte de ingresos de mercadería (recepciones de compra) · {escape(rango_txt)}{(' · ' + escape(filtros)) if filtros else ''}", e["sub"]),
             Spacer(1, 8),
             _resumen([("Recepciones", f"{t['recepciones']}"), ("Unidades que entraron", f"{t['unidades']}"),
                       ("Mercadería (sin ISV)", _m(t["subtotal"], moneda)), ("ISV", _m(t["isv"], moneda)),
                       ("Total comprado", _m(t["total"], moneda)), ("Saldo por pagar", _m(t["saldo"], moneda))], moneda)]

    def grupo(titulo, filas, col):
        if not filas:
            return
        story.append(Paragraph(titulo, e["seccion"]))
        datos_t = [[col, "Recepciones", "Unidades", "Mercadería", "ISV", "Total", "Pagado", "Saldo"]]
        for g in filas:
            datos_t.append([Paragraph(escape(str(g["clave"])), e["celda"]), g["recepciones"], g["unidades"], _m(g["subtotal"], moneda),
                            _m(g["isv"], moneda), _m(g["total"], moneda), _m(g["pagado"], moneda), _m(g["saldo"], moneda)])
        story.append(_tabla(datos_t, [6.6 * cm, 2.2 * cm, 2.0 * cm, 3.0 * cm, 2.6 * cm, 3.0 * cm, 2.9 * cm, 2.9 * cm], num_cols=(1, 2, 3, 4, 5, 6, 7)))

    grupo("Por proveedor", datos["por_proveedor"], "Proveedor")
    grupo("Por forma de pago", datos["por_forma"], "Forma de pago")

    story.append(Paragraph("Recepciones", e["seccion"]))
    datos_t = [["Fecha", "Recepción", "Solicitud", "Proveedor", "Factura", "Pago", "Unid.", "Mercadería", "ISV", "Total", "Saldo"]]
    for x in datos["recepciones"]:
        r = x["rec"]
        datos_t.append([_f(r.fecha, "%d/%m/%Y"), r.numero, r.solicitud.numero if r.solicitud else "", Paragraph(escape(r.proveedor or ""), e["celda"]),
                        Paragraph(escape(r.factura_proveedor or "—"), e["celda"]), r.forma_pago, x["unidades"], _m(r.subtotal, moneda),
                        _m(r.isv, moneda), _m(r.total, moneda), _m(r.saldo, moneda)])
    if len(datos_t) == 1:
        datos_t.append(["", "", "", "Sin recepciones en el período.", "", "", "", "", "", "", ""])
    else:
        datos_t.append(["", "", "", "Total", "", "", t["unidades"], _m(t["subtotal"], moneda), _m(t["isv"], moneda), _m(t["total"], moneda), _m(t["saldo"], moneda)])
    story.append(_tabla(datos_t, [2.0 * cm, 2.1 * cm, 2.1 * cm, 4.3 * cm, 2.4 * cm, 1.6 * cm, 1.2 * cm, 2.6 * cm, 2.2 * cm, 2.6 * cm, 2.4 * cm],
                        num_cols=(6, 7, 8, 9, 10), negrita_ultima=len(datos["recepciones"]) > 0))

    if datos["por_producto"]:
        story.append(Paragraph("Por producto", e["seccion"]))
        datos_t = [["Código", "Producto", "Categoría", "Unidades", "Costo promedio", "Valor (sin ISV)"]]
        for g in datos["por_producto"]:
            datos_t.append([Paragraph(escape(g["codigo"]), e["celda"]), Paragraph(escape(g["nombre"]), e["celda"]),
                            Paragraph(escape(g["categoria"] or "—"), e["celda"]), g["unidades"],
                            _m(g["costo_promedio"], moneda), _m(g["valor"], moneda)])
        story.append(_tabla(datos_t, [3.0 * cm, 9.0 * cm, 4.2 * cm, 2.2 * cm, 3.2 * cm, 3.4 * cm], num_cols=(3, 4, 5)))

    story.append(Paragraph("Detalle de productos recibidos", e["seccion"]))
    datos_t = [["Fecha", "Recepción", "Proveedor", "Producto", "Cant.", "Costo c/u", "Subtotal", "Había", "Quedó"]]
    for f in datos["filas"]:
        producto = f"{f['codigo']} · {f['nombre']}" + (f" ({f['categoria']})" if f["categoria"] else "")
        datos_t.append([_f(f["fecha"]), f["rec"].numero, Paragraph(escape(f["rec"].proveedor or ""), e["celda"]), Paragraph(escape(producto), e["celda"]),
                        f["cantidad"], _m(f["costo"], moneda), _m(f["subtotal"], moneda), f["antes"], f["despues"]])
    if len(datos_t) == 1:
        datos_t.append(["", "", "", "Sin productos recibidos en el período.", "", "", "", "", ""])
    story.append(_tabla(datos_t, [3.2 * cm, 2.1 * cm, 3.6 * cm, 8.0 * cm, 1.3 * cm, 2.4 * cm, 2.6 * cm, 1.3 * cm, 1.3 * cm], num_cols=(4, 5, 6, 7, 8)))
    story += [Spacer(1, 6), Paragraph("Costos sin ISV. «Había» y «Quedó»: existencia del producto justo antes y después de cada recepción.", e["sub"])]
    doc.build(story)
    return buffer.getvalue()


def excel_reporte_recepciones(datos: dict, rango_txt: str, nombre_taller: str, moneda: str, filtros: str = "") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor="1D4ED8")

    def encabezado(hoja, columnas):
        hoja.append(columnas)
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno

    def dinero(hoja, columnas):
        for col in columnas:
            hoja.cell(row=hoja.max_row, column=col).number_format = '#,##0.00'

    t = datos["total"]
    resumen = wb.active
    resumen.title = "Resumen"
    resumen.append([nombre_taller])
    resumen["A1"].font = Font(bold=True, size=14)
    resumen.append([f"Reporte de ingresos de mercadería — {rango_txt}" + (f" · {filtros}" if filtros else "")])
    resumen["A2"].font = Font(bold=True, size=12)
    for etiqueta, valor in (("Recepciones", t["recepciones"]), ("Unidades que entraron", t["unidades"]),
                            (f"Mercadería sin ISV ({moneda})", t["subtotal"]), (f"ISV ({moneda})", t["isv"]),
                            (f"Total comprado ({moneda})", t["total"]), (f"Pagado ({moneda})", t["pagado"]),
                            (f"Saldo por pagar ({moneda})", t["saldo"])):
        resumen.append([etiqueta, float(valor) if isinstance(valor, Decimal) else valor])
        if isinstance(valor, Decimal):
            dinero(resumen, (2,))
    for titulo, filas in (("Por proveedor", datos["por_proveedor"]), ("Por forma de pago", datos["por_forma"])):
        resumen.append([])
        resumen.append([titulo])
        resumen.cell(row=resumen.max_row, column=1).font = negrita
        encabezado(resumen, [titulo.replace("Por ", "").capitalize(), "Recepciones", "Unidades", f"Mercadería ({moneda})",
                             f"ISV ({moneda})", f"Total ({moneda})", f"Pagado ({moneda})", f"Saldo ({moneda})"])
        for g in filas:
            resumen.append([g["clave"], g["recepciones"], g["unidades"], float(g["subtotal"]), float(g["isv"]),
                            float(g["total"]), float(g["pagado"]), float(g["saldo"])])
            dinero(resumen, (4, 5, 6, 7, 8))

    hoja_r = wb.create_sheet("Recepciones")
    encabezado(hoja_r, ["Fecha", "Recepción", "Solicitud", "Proveedor", "Factura del proveedor", "Forma de pago", "Recibido por",
                        "Unidades", f"Mercadería ({moneda})", f"ISV ({moneda})", f"Total ({moneda})", f"Pagado ({moneda})", f"Saldo ({moneda})"])
    for x in datos["recepciones"]:
        r = x["rec"]
        hoja_r.append([_f(r.fecha), r.numero, r.solicitud.numero if r.solicitud else "", r.proveedor, r.factura_proveedor, r.forma_pago,
                       r.usuario_nombre, x["unidades"], float(_d(r.subtotal)), float(_d(r.isv)), float(_d(r.total)),
                       float(_d(r.pagado)), float(_d(r.saldo))])
        dinero(hoja_r, (9, 10, 11, 12, 13))

    hoja_p = wb.create_sheet("Por producto")
    encabezado(hoja_p, ["Código", "Producto", "Categoría", "Unidades", f"Costo promedio ({moneda})", f"Valor sin ISV ({moneda})"])
    for g in datos["por_producto"]:
        hoja_p.append([g["codigo"], g["nombre"], g["categoria"], g["unidades"], float(g["costo_promedio"]), float(g["valor"])])
        dinero(hoja_p, (5, 6))

    hoja_d = wb.create_sheet("Detalle")
    encabezado(hoja_d, ["Fecha", "Recepción", "Proveedor", "Factura", "Código", "Producto", "Categoría", "Cantidad",
                        f"Costo c/u ({moneda})", f"Subtotal ({moneda})", "Había", "Quedó"])
    for f in datos["filas"]:
        hoja_d.append([_f(f["fecha"]), f["rec"].numero, f["rec"].proveedor, f["rec"].factura_proveedor, f["codigo"], f["nombre"],
                       f["categoria"], f["cantidad"], float(f["costo"]), float(f["subtotal"]), f["antes"], f["despues"]])
        dinero(hoja_d, (9, 10))

    for hoja in (resumen, hoja_r, hoja_p, hoja_d):
        hoja.column_dimensions["A"].width = 28
        for col in "BCDEFGHIJKLM":
            hoja.column_dimensions[col].width = 16
    hoja_p.column_dimensions["B"].width = 42
    hoja_d.column_dimensions["F"].width = 42
    hoja_r.column_dimensions["D"].width = 28
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()
