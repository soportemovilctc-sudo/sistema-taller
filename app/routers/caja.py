"""Módulo de Caja: apertura y cierre de caja, conciliación bancaria y
egresos mayores, con dos vistas separadas por rol:

- Vista Cajera (rol "vendedor", ya etiquetado "Caja" en el resto del
  sistema): registra cobros del día (Efectivo, Transferencia, Tarjeta POS) y
  ve el cierre/arqueo diario de su caja chica. RESTRICCIÓN ABSOLUTA: no ve
  saldos de banco, liquidez consolidada ni reportes de utilidad.
- Vista Administrador (rol "admin"): saldos iniciales del mes, Liquidez
  Total, Estado de Resultados (Utilidad Neta), conciliación de las
  transferencias que registra la Cajera, y Gestión de Egresos Mayores.

Ambas vistas leen y escriben el mismo MovimientoFinanciero que ya usan
Contabilidad/Reportes/Dashboard (ver app/models.py), así que todo lo que
pasa por Caja queda automáticamente reflejado ahí también.
"""
from datetime import datetime, date, timedelta
from decimal import Decimal

from fastapi import APIRouter, Request, Depends, Form, UploadFile, File
from fastapi.responses import RedirectResponse, StreamingResponse, Response
from sqlalchemy.orm import Session
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from app.templates_env import templates
from app.database import get_db
from app.models import MovimientoFinanciero, Configuracion, METODOS_PAGO_CAJA, CATEGORIAS_EGRESO_MAYOR
from app.utils.calculations import to_decimal
from app.utils.caja_calculos import (
    calcular_liquidez, calcular_utilidad_neta, calcular_cierre_dia,
    guardar_saldo_inicial, mes_actual, CUENTA_CAJA_CHICA, CUENTA_BANCO,
)
from app.utils.flash import flash
from app.utils.pdf import generar_pdf_informe_caja
from app.deps import roles_required

router = APIRouter()


def _rango_fechas(periodo: str, desde: str, hasta: str):
    hoy = date.today()
    if periodo == "dia":
        return datetime.combine(hoy, datetime.min.time()), datetime.combine(hoy, datetime.max.time())
    if periodo == "semana":
        inicio = hoy - timedelta(days=hoy.weekday())
        return datetime.combine(inicio, datetime.min.time()), datetime.combine(hoy, datetime.max.time())
    if periodo == "rango" and desde and hasta:
        return (datetime.strptime(desde, "%Y-%m-%d"),
                datetime.combine(datetime.strptime(hasta, "%Y-%m-%d").date(), datetime.max.time()))
    # por defecto y para "mes": mes actual
    inicio = hoy.replace(day=1)
    return datetime.combine(inicio, datetime.min.time()), datetime.combine(hoy, datetime.max.time())


# ---------------------------------------------------------------------------
# Entrada: cada rol cae directo en su pantalla principal.
# ---------------------------------------------------------------------------
@router.get("/caja")
def caja_index(usuario=Depends(roles_required("vendedor"))):
    if usuario.get("rol") == "admin":
        return RedirectResponse("/caja/panel", status_code=303)
    return RedirectResponse("/caja/registrar", status_code=303)


# ---------------------------------------------------------------------------
# Vista Cajera: registrar cobro
# ---------------------------------------------------------------------------
@router.get("/caja/registrar")
def caja_registrar_form(request: Request, db: Session = Depends(get_db), usuario=Depends(roles_required("vendedor"))):
    hoy = date.today()
    desde = datetime.combine(hoy, datetime.min.time())
    hasta = datetime.combine(hoy, datetime.max.time())
    transacciones_hoy = db.query(MovimientoFinanciero).filter(
        MovimientoFinanciero.metodo_pago.isnot(None),
        MovimientoFinanciero.fecha.between(desde, hasta),
    ).order_by(MovimientoFinanciero.fecha.desc()).all()

    return templates.TemplateResponse("caja/registrar.html", {
        "request": request, "usuario": usuario,
        "metodos_pago": METODOS_PAGO_CAJA,
        "transacciones_hoy": transacciones_hoy,
    })


@router.post("/caja/registrar")
def caja_registrar_crear(
    request: Request,
    monto: str = Form(...), cliente_nombre: str = Form(""), metodo_pago: str = Form(...),
    num_referencia: str = Form(""), comprobante: UploadFile | None = File(None),
    db: Session = Depends(get_db), usuario=Depends(roles_required("vendedor")),
):
    if metodo_pago not in METODOS_PAGO_CAJA:
        flash(request, "Método de pago no válido.", "error")
        return RedirectResponse("/caja/registrar", status_code=303)

    try:
        monto_dec = to_decimal(monto)
        if monto_dec <= 0:
            raise ValueError("El monto debe ser mayor a cero.")
    except ValueError as e:
        flash(request, str(e), "error")
        return RedirectResponse("/caja/registrar", status_code=303)

    if metodo_pago == "Transferencia" and not num_referencia.strip():
        flash(request, "El número de comprobante/referencia es obligatorio para una transferencia.", "error")
        return RedirectResponse("/caja/registrar", status_code=303)

    # Efectivo -> caja chica; Tarjeta (POS) y Transferencia -> banco (ahí es
    # donde realmente cae el dinero, no en la caja física).
    cuenta = CUENTA_CAJA_CHICA if metodo_pago == "Efectivo" else CUENTA_BANCO

    comprobante_data = None
    comprobante_mime = None
    if comprobante and comprobante.filename:
        contenido = comprobante.file.read()
        if contenido:
            comprobante_data = contenido
            comprobante_mime = comprobante.content_type or "application/octet-stream"

    mov = MovimientoFinanciero(
        tipo="ingreso", categoria="Cobro de caja", monto=monto_dec.quantize(Decimal("0.01")),
        descripcion=f"Cobro en {metodo_pago}" + (f" - {cliente_nombre.strip()}" if cliente_nombre.strip() else ""),
        usuario_nombre=usuario["nombre_completo"], referencia="Caja",
        cuenta=cuenta, metodo_pago=metodo_pago,
        cliente_nombre=cliente_nombre.strip() or None,
        num_referencia=num_referencia.strip() or None,
        comprobante_data=comprobante_data, comprobante_mime=comprobante_mime,
        estado_conciliacion="pendiente" if metodo_pago == "Transferencia" else None,
    )
    db.add(mov)
    db.commit()
    flash(request, "Cobro registrado correctamente.", "success")
    return RedirectResponse("/caja/registrar", status_code=303)


@router.get("/caja/comprobante/{mov_id}")
def caja_ver_comprobante(mov_id: int, db: Session = Depends(get_db), usuario=Depends(roles_required("vendedor"))):
    mov = db.get(MovimientoFinanciero, mov_id)
    if not mov or not mov.comprobante_data:
        return Response(status_code=404)
    return StreamingResponse(BytesIO(mov.comprobante_data), media_type=mov.comprobante_mime or "application/octet-stream")


# ---------------------------------------------------------------------------
# Vista Cajera: cierre / arqueo diario (solo lectura, se calcula al vuelo)
# ---------------------------------------------------------------------------
@router.get("/caja/cierre")
def caja_cierre(request: Request, db: Session = Depends(get_db), usuario=Depends(roles_required("vendedor"))):
    cierre = calcular_cierre_dia(db)
    return templates.TemplateResponse("caja/cierre.html", {
        "request": request, "usuario": usuario, "cierre": cierre,
    })


# ---------------------------------------------------------------------------
# Vista Administrador: panel gerencial (liquidez + rentabilidad + saldos)
# ---------------------------------------------------------------------------
@router.get("/caja/panel")
def caja_panel(
    request: Request, periodo: str = "mes", desde: str = "", hasta: str = "",
    db: Session = Depends(get_db), usuario=Depends(roles_required("admin")),
):
    anio, mes = mes_actual()
    liquidez = calcular_liquidez(db, anio, mes)
    ini, fin = _rango_fechas(periodo, desde, hasta)
    utilidad = calcular_utilidad_neta(db, ini, fin)

    return templates.TemplateResponse("caja/panel.html", {
        "request": request, "usuario": usuario,
        "liquidez": liquidez, "utilidad": utilidad,
        "periodo": periodo, "desde": desde, "hasta": hasta,
        "anio": anio, "mes": mes,
    })


@router.post("/caja/panel/saldos")
def caja_panel_guardar_saldos(
    request: Request, anio: int = Form(...), mes: int = Form(...),
    saldo_inicial_caja_chica: str = Form("0"), saldo_inicial_banco: str = Form("0"),
    db: Session = Depends(get_db), usuario=Depends(roles_required("admin")),
):
    guardar_saldo_inicial(db, anio, mes, saldo_inicial_caja_chica, saldo_inicial_banco, usuario["nombre_completo"])
    db.commit()
    flash(request, "Saldos iniciales del mes guardados correctamente.", "success")
    return RedirectResponse("/caja/panel", status_code=303)


# ---------------------------------------------------------------------------
# Vista Administrador: conciliación bancaria + egresos mayores
# ---------------------------------------------------------------------------
@router.get("/caja/conciliacion")
def caja_conciliacion(request: Request, db: Session = Depends(get_db), usuario=Depends(roles_required("admin"))):
    pendientes = db.query(MovimientoFinanciero).filter(
        MovimientoFinanciero.metodo_pago == "Transferencia",
        MovimientoFinanciero.estado_conciliacion == "pendiente",
    ).order_by(MovimientoFinanciero.fecha.asc()).all()
    conciliadas_recientes = db.query(MovimientoFinanciero).filter(
        MovimientoFinanciero.metodo_pago == "Transferencia",
        MovimientoFinanciero.estado_conciliacion == "conciliado",
    ).order_by(MovimientoFinanciero.conciliado_en.desc()).limit(20).all()

    egresos = db.query(MovimientoFinanciero).filter(
        MovimientoFinanciero.tipo == "gasto", MovimientoFinanciero.cuenta == CUENTA_BANCO,
    ).order_by(MovimientoFinanciero.fecha.desc()).limit(50).all()
    total_egresos = sum((to_decimal(m.monto) for m in egresos), Decimal("0.00"))

    return templates.TemplateResponse("caja/conciliacion.html", {
        "request": request, "usuario": usuario,
        "pendientes": pendientes, "conciliadas_recientes": conciliadas_recientes,
        "egresos": egresos, "total_egresos": total_egresos,
        "categorias_egreso": CATEGORIAS_EGRESO_MAYOR,
    })


@router.post("/caja/conciliacion/{mov_id}/validar")
def caja_conciliacion_validar(mov_id: int, request: Request, db: Session = Depends(get_db), usuario=Depends(roles_required("admin"))):
    mov = db.get(MovimientoFinanciero, mov_id)
    if not mov or mov.metodo_pago != "Transferencia":
        flash(request, "Transferencia no encontrada.", "error")
        return RedirectResponse("/caja/conciliacion", status_code=303)

    mov.estado_conciliacion = "conciliado"
    mov.conciliado_por = usuario["nombre_completo"]
    mov.conciliado_en = datetime.utcnow()
    db.commit()
    flash(request, "Transferencia marcada como conciliada.", "success")
    return RedirectResponse("/caja/conciliacion", status_code=303)


@router.post("/caja/egresos/nuevo")
def caja_egreso_crear(
    request: Request, categoria: str = Form(...), monto: str = Form(...), descripcion: str = Form(""),
    db: Session = Depends(get_db), usuario=Depends(roles_required("admin")),
):
    try:
        monto_dec = to_decimal(monto)
        if monto_dec <= 0:
            raise ValueError("El monto debe ser mayor a cero.")
    except ValueError as e:
        flash(request, str(e), "error")
        return RedirectResponse("/caja/conciliacion", status_code=303)

    categoria_limpia = categoria.strip() or "Otros"
    db.add(MovimientoFinanciero(
        tipo="gasto", categoria=categoria_limpia, monto=monto_dec.quantize(Decimal("0.01")),
        descripcion=descripcion.strip(), usuario_nombre=usuario["nombre_completo"],
        referencia="Egreso mayor", cuenta=CUENTA_BANCO,
    ))
    db.commit()
    flash(request, "Egreso mayor registrado correctamente.", "success")
    return RedirectResponse("/caja/conciliacion", status_code=303)


# ---------------------------------------------------------------------------
# Vista Administrador: informe mensual (con un clic, PDF y Excel)
# ---------------------------------------------------------------------------
MESES_NOMBRE = ["", "Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio", "Julio",
                "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]


def _contexto_informe(db: Session, anio: int, mes: int) -> dict:
    """Arma todos los datos del informe mensual de Caja: liquidez, utilidad
    neta, transferencias y egresos mayores del mes indicado."""
    desde = datetime(anio, mes, 1)
    hasta = (datetime(anio + 1, 1, 1) if mes == 12 else datetime(anio, mes + 1, 1)) - timedelta(seconds=1)

    liquidez = calcular_liquidez(db, anio, mes)
    utilidad = calcular_utilidad_neta(db, desde, hasta)

    transferencias = db.query(MovimientoFinanciero).filter(
        MovimientoFinanciero.metodo_pago == "Transferencia",
        MovimientoFinanciero.fecha.between(desde, hasta),
    ).order_by(MovimientoFinanciero.fecha.asc()).all()

    egresos_mayores = db.query(MovimientoFinanciero).filter(
        MovimientoFinanciero.tipo == "gasto", MovimientoFinanciero.cuenta == CUENTA_BANCO,
        MovimientoFinanciero.fecha.between(desde, hasta),
    ).order_by(MovimientoFinanciero.fecha.asc()).all()

    cfg = db.get(Configuracion, 1)
    return {
        "anio": anio, "mes": mes, "nombre_mes": MESES_NOMBRE[mes],
        "nombre_taller": cfg.nombre_taller if cfg else "Mi Taller",
        "moneda": cfg.moneda if cfg else "L",
        "liquidez": liquidez, "utilidad": utilidad,
        "transferencias": transferencias, "egresos_mayores": egresos_mayores,
    }


def _anio_mes_desde_query(anio: int, mes: int):
    if not anio or not mes:
        return mes_actual()
    return anio, mes


@router.get("/caja/informe")
def caja_informe(
    request: Request, anio: int = 0, mes: int = 0,
    db: Session = Depends(get_db), usuario=Depends(roles_required("admin")),
):
    anio, mes = _anio_mes_desde_query(anio, mes)
    ctx = _contexto_informe(db, anio, mes)
    return templates.TemplateResponse("caja/informe.html", {"request": request, "usuario": usuario, **ctx})


@router.get("/caja/informe/exportar/pdf")
def caja_informe_pdf(anio: int = 0, mes: int = 0, db: Session = Depends(get_db), usuario=Depends(roles_required("admin"))):
    anio, mes = _anio_mes_desde_query(anio, mes)
    ctx = _contexto_informe(db, anio, mes)

    pdf_bytes = generar_pdf_informe_caja({
        "nombre_taller": ctx["nombre_taller"], "moneda": ctx["moneda"],
        "subtitulo": f"{ctx['nombre_mes']} {ctx['anio']}",
        "liquidez": ctx["liquidez"], "utilidad": ctx["utilidad"],
        "transferencias": [
            {
                "fecha": t.fecha.strftime("%d/%m/%Y"), "cliente": t.cliente_nombre,
                "referencia": t.num_referencia, "monto": t.monto,
                "estado": "Conciliada" if t.estado_conciliacion == "conciliado" else "Pendiente",
            }
            for t in ctx["transferencias"]
        ],
        "egresos_mayores": [
            {"fecha": e.fecha.strftime("%d/%m/%Y"), "descripcion": e.descripcion, "categoria": e.categoria,
             "monto": e.monto, "usuario": e.usuario_nombre}
            for e in ctx["egresos_mayores"]
        ],
    })
    nombre_archivo = f"informe_caja_{ctx['anio']}_{ctx['mes']:02d}.pdf"
    return StreamingResponse(BytesIO(pdf_bytes), media_type="application/pdf",
                              headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'})


@router.get("/caja/informe/exportar/excel")
def caja_informe_excel(anio: int = 0, mes: int = 0, db: Session = Depends(get_db), usuario=Depends(roles_required("admin"))):
    anio, mes = _anio_mes_desde_query(anio, mes)
    ctx = _contexto_informe(db, anio, mes)
    liq, util = ctx["liquidez"], ctx["utilidad"]

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Resumen"
    titulo_font = Font(name="Calibri", size=14, bold=True)
    hoja.merge_cells("A1:B1")
    hoja["A1"] = ctx["nombre_taller"]
    hoja["A1"].font = titulo_font
    hoja.merge_cells("A2:B2")
    hoja["A2"] = f"Informe mensual de Caja — {ctx['nombre_mes']} {ctx['anio']}"
    hoja["A2"].font = Font(name="Calibri", size=12, bold=True)

    fill = PatternFill(start_color="3D5F96", end_color="3D5F96", fill_type="solid")
    negrita = Font(bold=True)

    hoja["A4"] = "Liquidez Total (saldo consolidado del mes)"
    hoja["A4"].font = Font(bold=True, color="FFFFFF")
    hoja["A4"].fill = fill
    hoja.merge_cells("A4:B4")
    filas_liq = [
        ("Saldo inicial Caja Chica", float(liq["saldo_inicial_caja_chica"])),
        ("Saldo inicial Banco", float(liq["saldo_inicial_banco"])),
        ("(+) Ingresos Caja Chica", float(liq["ingresos_caja_chica"])),
        ("(-) Egresos Caja Chica", float(liq["egresos_caja_chica"])),
        ("(+) Ingresos Banco", float(liq["ingresos_banco"])),
        ("(-) Egresos Banco", float(liq["egresos_banco"])),
        ("Caja Chica actual", float(liq["caja_chica_actual"])),
        ("Banco actual", float(liq["banco_actual"])),
        ("LIQUIDEZ TOTAL", float(liq["liquidez_total"])),
    ]
    fila = 5
    for etiqueta, valor in filas_liq:
        hoja.cell(row=fila, column=1, value=etiqueta)
        celda_valor = hoja.cell(row=fila, column=2, value=valor)
        celda_valor.number_format = "#,##0.00"
        if etiqueta == "LIQUIDEZ TOTAL":
            hoja.cell(row=fila, column=1).font = negrita
            celda_valor.font = negrita
        fila += 1

    fila += 1
    hoja.cell(row=fila, column=1, value="Estado de Resultados (rentabilidad del período)")
    hoja.cell(row=fila, column=1).font = Font(bold=True, color="FFFFFF")
    hoja.cell(row=fila, column=1).fill = fill
    hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=2)
    fila += 1
    filas_util = [
        ("Ingresos del período", float(util["ingresos"])),
        ("Egresos / gastos del período", float(util["egresos"])),
        ("UTILIDAD NETA", float(util["utilidad_neta"])),
    ]
    for etiqueta, valor in filas_util:
        hoja.cell(row=fila, column=1, value=etiqueta)
        celda_valor = hoja.cell(row=fila, column=2, value=valor)
        celda_valor.number_format = "#,##0.00"
        if etiqueta == "UTILIDAD NETA":
            hoja.cell(row=fila, column=1).font = negrita
            celda_valor.font = negrita
        fila += 1
    hoja.column_dimensions["A"].width = 34
    hoja.column_dimensions["B"].width = 18

    def _hoja_detalle(nombre, encabezados, filas, columnas_totales=None):
        hoja2 = wb.create_sheet(nombre[:31])
        borde = Border(left=Side(style="thin"), right=Side(style="thin"), top=Side(style="thin"), bottom=Side(style="thin"))
        for col, titulo in enumerate(encabezados, start=1):
            celda = hoja2.cell(row=1, column=col, value=titulo)
            celda.fill = fill
            celda.font = Font(bold=True, color="FFFFFF")
            celda.border = borde
        for i, f in enumerate(filas, start=2):
            for col, valor in enumerate(f, start=1):
                celda = hoja2.cell(row=i, column=col, value=valor)
                celda.border = borde
                if isinstance(valor, float):
                    celda.number_format = "#,##0.00"
        for col_idx, titulo in enumerate(encabezados, start=1):
            largo = len(str(titulo))
            for f in filas:
                if col_idx - 1 < len(f):
                    largo = max(largo, len(str(f[col_idx - 1])))
            hoja2.column_dimensions[get_column_letter(col_idx)].width = min(max(largo + 2, 10), 40)
        return hoja2

    _hoja_detalle(
        "Transferencias",
        ["Fecha", "Cliente", "Referencia", "Monto", "Estado"],
        [[t.fecha.strftime("%d/%m/%Y"), t.cliente_nombre or "-", t.num_referencia or "-", float(t.monto),
          "Conciliada" if t.estado_conciliacion == "conciliado" else "Pendiente"] for t in ctx["transferencias"]],
    )
    _hoja_detalle(
        "Egresos mayores",
        ["Fecha", "Concepto", "Categoría", "Monto", "Usuario"],
        [[e.fecha.strftime("%d/%m/%Y"), e.descripcion or "-", e.categoria, float(e.monto), e.usuario_nombre or "-"]
         for e in ctx["egresos_mayores"]],
    )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    nombre_archivo = f"informe_caja_{ctx['anio']}_{ctx['mes']:02d}.xlsx"
    return StreamingResponse(
        buffer, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre_archivo}"'},
    )
