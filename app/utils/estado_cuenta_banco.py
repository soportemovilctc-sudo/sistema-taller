"""Estado de cuenta bancario: todo lo que entró y salió de la cuenta Banco en
un período, con el saldo después de cada movimiento.

Usa los mismos registros y el mismo saldo que «Banco actual» del Panel
gerencial (MovimientoFinanciero con cuenta = banco): un abono de orden por
transferencia o tarjeta aparece el día en que se facturó la orden.

Tipos de movimiento:
  * Transferencia   — ingresos por transferencia (facturas, cobros de Caja...).
  * Tarjeta / POS   — ingresos con tarjeta (ventas del POS, facturas, cobros).
  * Depósito        — efectivo de Caja Chica depositado en el banco.
  * Otro ingreso    — otra entrada al banco (sin forma de pago registrada).
  * Salida          — egresos mayores, compras y otros pagos hechos por banco.
"""
from datetime import date, datetime, timedelta
from decimal import Decimal
from io import BytesIO

from sqlalchemy.orm import Session

from app.models import MovimientoFinanciero, CATEGORIA_DEPOSITO_BANCO
from app.utils.calculations import to_decimal
from app.utils.caja_calculos import CUENTA_BANCO
from app.utils.depositos import es_convertible
from app.utils.turnos_caja import inicio_dia_local_en_utc, a_hora_local, hoy_local

CERO = Decimal("0.00")
TIPOS = {
    "transferencia": "Transferencia",
    "tarjeta": "Tarjeta / POS",
    "deposito": "Depósito",
    "otro_ingreso": "Otro ingreso",
    "salida": "Salida",
}


def q(v) -> Decimal:
    return to_decimal(v or 0).quantize(Decimal("0.01"))


def clasificar(m: MovimientoFinanciero) -> str:
    if m.tipo != "ingreso":
        return "salida"
    if m.categoria == CATEGORIA_DEPOSITO_BANCO:
        return "deposito"
    metodo = (m.metodo_pago or "").lower()
    if "tarjeta" in metodo:
        return "tarjeta"
    if "transfer" in metodo:
        return "transferencia"
    return "otro_ingreso"


def _detalle(m: MovimientoFinanciero) -> str:
    partes = [m.categoria or ""]
    texto = (m.cliente_nombre or "").strip() or (m.descripcion or "").strip()
    if texto and texto != m.categoria:
        partes.append(texto)
    return " · ".join(p for p in partes if p)


def _referencia(m: MovimientoFinanciero) -> str:
    refs = [r for r in ((m.num_referencia or "").strip(), (m.referencia or "").strip()) if r]
    return " · ".join(dict.fromkeys(refs))


def _conciliacion(m: MovimientoFinanciero) -> str:
    if m.estado_conciliacion == "conciliado":
        return "Conciliada"
    if m.estado_conciliacion == "pendiente":
        return "Pendiente"
    return ""


def periodo_desde_texto(desde_txt: str, hasta_txt: str) -> tuple[date, date]:
    def leer(t):
        try:
            return datetime.strptime((t or "").strip(), "%Y-%m-%d").date()
        except ValueError:
            return None
    hoy = hoy_local()
    desde, hasta = leer(desde_txt), leer(hasta_txt)
    if not desde and not hasta:
        return hoy.replace(day=1), hoy
    desde = desde or (hasta or hoy).replace(day=1)
    hasta = hasta or hoy
    if hasta < desde:
        desde, hasta = hasta, desde
    return desde, hasta


def estado_cuenta(db: Session, desde: date, hasta: date, tipo: str = "") -> dict:
    from app.utils.estados_financieros import saldo_efectivo_a  # evita import circular

    inicio = inicio_dia_local_en_utc(desde)
    fin = inicio_dia_local_en_utc(hasta + timedelta(days=1))
    saldo_inicial = q(saldo_efectivo_a(db, inicio)["banco"])

    movs = (
        db.query(MovimientoFinanciero)
        .filter(MovimientoFinanciero.cuenta == CUENTA_BANCO,
                MovimientoFinanciero.excluir_de_contabilidad == False,  # noqa: E712
                MovimientoFinanciero.fecha >= inicio, MovimientoFinanciero.fecha < fin)
        .order_by(MovimientoFinanciero.fecha.asc(), MovimientoFinanciero.id.asc())
        .all()
    )
    # Inicios de mes dentro del período: si ese mes tiene un saldo inicial
    # configurado en Caja distinto de lo que venía, se muestra el ajuste
    # (así el saldo final cuadra con «Banco actual» del Panel gerencial).
    cortes = []
    d = (desde.replace(day=1) + timedelta(days=32)).replace(day=1)
    while d <= hasta:
        cortes.append(d)
        d = (d + timedelta(days=32)).replace(day=1)

    filas, saldo = [], saldo_inicial
    totales = {k: {"n": 0, "monto": CERO} for k in TIPOS}
    ajustes = CERO

    def aplicar_corte(dia):
        nonlocal saldo, ajustes
        esperado = q(saldo_efectivo_a(db, inicio_dia_local_en_utc(dia))["banco"])
        if esperado != saldo:
            dif = q(esperado - saldo)
            ajustes += dif
            saldo = esperado
            filas.append({"fecha": None, "dia": dia, "tipo": "ajuste", "tipo_texto": "Ajuste de saldo",
                          "detalle": f"Saldo inicial de Banco configurado para {dia.strftime('%m/%Y')} en Caja → Panel gerencial",
                          "referencia": "", "entrada": dif if dif > 0 else CERO, "salida": -dif if dif < 0 else CERO,
                          "saldo": saldo, "conciliacion": "", "id": None})

    i_corte = 0
    for m in movs:
        while i_corte < len(cortes) and m.fecha >= inicio_dia_local_en_utc(cortes[i_corte]):
            aplicar_corte(cortes[i_corte])
            i_corte += 1
        clase = clasificar(m)
        monto = q(m.monto)
        saldo = q(saldo + (monto if m.tipo == "ingreso" else -monto))
        totales[clase]["n"] += 1
        totales[clase]["monto"] += monto
        filas.append({
            "fecha": m.fecha, "dia": a_hora_local(m.fecha).date(), "tipo": clase, "tipo_texto": TIPOS[clase],
            "detalle": _detalle(m), "referencia": _referencia(m),
            "entrada": monto if m.tipo == "ingreso" else CERO, "salida": monto if m.tipo != "ingreso" else CERO,
            "saldo": saldo, "conciliacion": _conciliacion(m), "id": m.id, "usuario": m.usuario_nombre or "",
            "convertible": es_convertible(m),
        })
    while i_corte < len(cortes):
        aplicar_corte(cortes[i_corte])
        i_corte += 1

    entradas = q(sum((totales[k]["monto"] for k in TIPOS if k != "salida"), CERO))
    salidas = q(totales["salida"]["monto"])
    visibles = [f for f in filas if not tipo or f["tipo"] == tipo or (tipo == "entradas" and f["tipo"] not in ("salida", "ajuste"))]
    return {
        "desde": desde, "hasta": hasta, "tipo": tipo,
        "saldo_inicial": saldo_inicial, "saldo_final": saldo,
        "entradas": entradas, "salidas": salidas, "ajustes": q(ajustes),
        "totales": {k: {"n": v["n"], "monto": q(v["monto"])} for k, v in totales.items()},
        "filas": visibles, "n_total": len(movs),
        "pendientes_conciliar": sum(1 for f in filas if f["conciliacion"] == "Pendiente"),
        "tipos": TIPOS,
    }


def _titulo_periodo(ec: dict) -> str:
    return f"Del {ec['desde'].strftime('%d/%m/%Y')} al {ec['hasta'].strftime('%d/%m/%Y')}"


def _fecha_txt(f) -> str:
    return a_hora_local(f["fecha"]).strftime("%d/%m/%Y %I:%M %p") if f["fecha"] else f["dia"].strftime("%d/%m/%Y")


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def excel_estado_cuenta(ec: dict, nombre_taller: str, moneda: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()
    hoja = wb.active
    hoja.title = "Estado de cuenta"
    hoja["A1"] = nombre_taller
    hoja["A1"].font = Font(size=14, bold=True)
    hoja["A2"] = f"Estado de cuenta bancario · {_titulo_periodo(ec)}"
    hoja["A2"].font = Font(size=12, bold=True)
    fila = 4
    resumen = [("Saldo inicial", ec["saldo_inicial"]), ("Entradas", ec["entradas"]), ("Salidas", ec["salidas"])]
    if ec["ajustes"]:
        resumen.append(("Ajustes de saldo inicial", ec["ajustes"]))
    resumen.append(("Saldo final", ec["saldo_final"]))
    for etiqueta, valor in resumen:
        hoja.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        c = hoja.cell(row=fila, column=2, value=float(valor))
        c.number_format = '#,##0.00'
        fila += 1
    fila += 1
    hoja.cell(row=fila, column=1, value="Por tipo").font = Font(bold=True)
    fila += 1
    for k, nombre in TIPOS.items():
        hoja.cell(row=fila, column=1, value=nombre)
        hoja.cell(row=fila, column=2, value=ec["totales"][k]["n"])
        c = hoja.cell(row=fila, column=3, value=float(ec["totales"][k]["monto"]))
        c.number_format = '#,##0.00'
        fila += 1
    fila += 1
    encabezados = ["Fecha", "Tipo", "Detalle", "Referencia", f"Entrada ({moneda})", f"Salida ({moneda})", f"Saldo ({moneda})", "Conciliación", "Usuario"]
    relleno = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    for col, t in enumerate(encabezados, start=1):
        c = hoja.cell(row=fila, column=col, value=t)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = relleno
        c.alignment = Alignment(horizontal="center")
    fila += 1
    for f in ec["filas"]:
        valores = [_fecha_txt(f), f["tipo_texto"], f["detalle"], f["referencia"],
                   float(f["entrada"]) if f["entrada"] else None, float(f["salida"]) if f["salida"] else None,
                   float(f["saldo"]), f["conciliacion"], f.get("usuario", "")]
        for col, v in enumerate(valores, start=1):
            c = hoja.cell(row=fila, column=col, value=v)
            if col in (5, 6, 7):
                c.number_format = '#,##0.00'
        fila += 1
    for col, ancho in zip("ABCDEFGHI", (20, 15, 46, 22, 14, 14, 14, 13, 16)):
        hoja.column_dimensions[col].width = ancho
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ---------------------------------------------------------------------------
# PDF
# ---------------------------------------------------------------------------
def pdf_estado_cuenta(ec: dict, nombre_taller: str, moneda: str) -> bytes:
    from xml.sax.saxutils import escape
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter, landscape
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

    def m(v):
        return f"{moneda} {to_decimal(v):,.2f}"

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), topMargin=1.1 * cm, bottomMargin=1.1 * cm,
                            leftMargin=1.2 * cm, rightMargin=1.2 * cm, title="Estado de cuenta bancario")
    estilos = getSampleStyleSheet()
    titulo = ParagraphStyle("t", parent=estilos["Title"], fontSize=15, alignment=0, spaceAfter=2)
    sub = ParagraphStyle("s", parent=estilos["Normal"], fontSize=9.5, textColor=colors.HexColor("#475569"))
    celda = ParagraphStyle("c", parent=estilos["Normal"], fontSize=7.8, leading=9.5)
    azul = colors.HexColor("#1d4ed8")
    story = [Paragraph(escape(nombre_taller), titulo),
             Paragraph(f"Estado de cuenta bancario · {_titulo_periodo(ec)}", sub), Spacer(1, 8)]

    resumen = [["Saldo inicial", "Entradas", "Salidas"] + (["Ajustes"] if ec["ajustes"] else []) + ["Saldo final"],
               [m(ec["saldo_inicial"]), m(ec["entradas"]), m(ec["salidas"])] + ([m(ec["ajustes"])] if ec["ajustes"] else []) + [m(ec["saldo_final"])]]
    t = Table(resumen)
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, 0), 8), ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#475569")),
        ("FONTNAME", (0, 1), (-1, 1), "Helvetica-Bold"), ("FONTSIZE", (0, 1), (-1, 1), 12),
        ("FONTSIZE", (-1, 1), (-1, 1), 14), ("TEXTCOLOR", (-1, 1), (-1, 1), azul),
        ("RIGHTPADDING", (0, 0), (-1, -1), 18), ("LEFTPADDING", (0, 0), (-1, -1), 0),
    ]))
    t.hAlign = "LEFT"
    story += [t, Spacer(1, 8)]

    por_tipo = [[f"{TIPOS[k]}: {ec['totales'][k]['n']} · {m(ec['totales'][k]['monto'])}" for k in TIPOS]]
    t = Table(por_tipo)
    t.setStyle(TableStyle([("FONTSIZE", (0, 0), (-1, -1), 8), ("TEXTCOLOR", (0, 0), (-1, -1), colors.HexColor("#334155")),
                           ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 14)]))
    t.hAlign = "LEFT"
    story += [t, Spacer(1, 10)]

    datos = [["Fecha", "Tipo", "Detalle", "Referencia", "Entrada", "Salida", "Saldo", "Conc."]]
    if not ec["tipo"]:
        datos.append([ec["desde"].strftime("%d/%m/%Y"), "", "Saldo inicial", "", "", "", m(ec["saldo_inicial"]), ""])
    for f in ec["filas"]:
        datos.append([_fecha_txt(f), f["tipo_texto"], Paragraph(escape(f["detalle"]), celda), Paragraph(escape(f["referencia"]), celda),
                      m(f["entrada"]) if f["entrada"] else "", m(f["salida"]) if f["salida"] else "", m(f["saldo"]), f["conciliacion"]])
    if not ec["filas"]:
        datos.append(["", "", "Sin movimientos en el período.", "", "", "", "", ""])
    t = Table(datos, colWidths=[3.2 * cm, 2.4 * cm, 7.2 * cm, 4.0 * cm, 2.4 * cm, 2.4 * cm, 2.5 * cm, 1.6 * cm], repeatRows=1)
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), azul), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 7.8),
        ("ALIGN", (4, 0), (6, -1), "RIGHT"), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#cbd5e1")),
        ("FONTNAME", (6, 1), (6, -1), "Helvetica-Bold"),
    ]))
    story.append(t)
    story += [Spacer(1, 8), Paragraph(
        "Movimientos de la cuenta Banco registrados en el sistema. Un abono por transferencia o tarjeta aparece el día en que se facturó la orden.", sub)]
    doc.build(story)
    return buffer.getvalue()
