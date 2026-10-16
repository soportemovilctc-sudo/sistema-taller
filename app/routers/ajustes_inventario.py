"""Ajustes de inventario (solo administrador): ajuste de varios productos a
la vez, importación desde Excel con vista previa, historial con búsqueda,
detalle, reversión y reportes en unidades y lempiras (pantalla y Excel).
La lógica está en app/utils/inventario_costos.py."""
import io
import json
from datetime import timedelta
from decimal import Decimal

from fastapi import APIRouter, Request, Depends, Form, UploadFile, File
from fastapi.responses import RedirectResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import or_
from sqlalchemy.orm import Session

from app.templates_env import templates
from app.database import get_db
from app.deps import roles_required
from app.models import (
    Producto, AjusteInventario, AjusteInventarioDetalle, Categoria,
    MOTIVOS_AJUSTE_INVENTARIO, MOTIVO_REVERSION_AJUSTE,
)
from app.utils.flash import flash
from app.utils.inventario_costos import (
    aplicar_ajuste, revertir_ajuste, calcular_lineas, reporte_ajustes, MOTIVO_AJUSTE_RAPIDO,
)
from app.utils.turnos_caja import a_hora_local, hoy_local, inicio_dia_local_en_utc

router = APIRouter()
BASE = "/ajustes-inventario"
solo_admin = roles_required("admin")


def _fecha(valor: str):
    from datetime import datetime
    try:
        return datetime.strptime((valor or "").strip(), "%Y-%m-%d").date()
    except ValueError:
        return None


def _rango(desde: str, hasta: str, por_defecto_mes: bool):
    """Fechas locales (Honduras) -> rango UTC [desde, hasta)."""
    d, h = _fecha(desde), _fecha(hasta)
    if por_defecto_mes and not d and not h:
        hoy = hoy_local()
        d, h = hoy.replace(day=1), hoy
    inicio = inicio_dia_local_en_utc(d) if d else None
    fin = inicio_dia_local_en_utc(h + timedelta(days=1)) if h else None
    return d, h, inicio, fin


def _productos_activos(db: Session):
    return db.query(Producto).filter(Producto.estado == "activo").order_by(Producto.categoria, Producto.nombre).all()


# ---------------------------------------------------------------------------
# Historial
# ---------------------------------------------------------------------------
@router.get(BASE)
def ajustes_lista(
    request: Request, desde: str = "", hasta: str = "", motivo: str = "", q: str = "",
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    d, h, inicio, fin = _rango(desde, hasta, por_defecto_mes=False)
    consulta = db.query(AjusteInventario)
    if inicio:
        consulta = consulta.filter(AjusteInventario.fecha >= inicio)
    if fin:
        consulta = consulta.filter(AjusteInventario.fecha < fin)
    if motivo:
        consulta = consulta.filter(AjusteInventario.motivo == motivo)
    if q.strip():
        like = f"%{q.strip()}%"
        consulta = consulta.filter(or_(
            AjusteInventario.numero.ilike(like),
            AjusteInventario.observacion.ilike(like),
            AjusteInventario.id.in_(
                db.query(AjusteInventarioDetalle.ajuste_id).filter(or_(
                    AjusteInventarioDetalle.producto_codigo.ilike(like),
                    AjusteInventarioDetalle.producto_nombre.ilike(like),
                ))
            ),
        ))
    hay_filtro = bool(inicio or fin or motivo or q.strip())
    ajustes = consulta.order_by(AjusteInventario.fecha.desc(), AjusteInventario.id.desc()).limit(500 if hay_filtro else 100).all()
    return templates.TemplateResponse("ajustes_inventario/lista.html", {
        "request": request, "usuario": usuario, "ajustes": ajustes,
        "desde": desde, "hasta": hasta, "motivo": motivo, "q": q, "hay_filtro": hay_filtro,
        "motivos": MOTIVOS_AJUSTE_INVENTARIO + [MOTIVO_REVERSION_AJUSTE],
        "hoy_iso": hoy_local().isoformat(), "a_hora_local": a_hora_local,
    })


# ---------------------------------------------------------------------------
# Nuevo ajuste (varios productos)
# ---------------------------------------------------------------------------
@router.get(BASE + "/nuevo")
def ajustes_nuevo_form(request: Request, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    return templates.TemplateResponse("ajustes_inventario/nuevo.html", {
        "request": request, "usuario": usuario, "productos": _productos_activos(db),
        "motivos": MOTIVOS_AJUSTE_INVENTARIO,
    })


@router.post(BASE + "/nuevo")
def ajustes_nuevo_crear(
    request: Request,
    motivo: str = Form(""), observacion: str = Form(""),
    producto_id: list[str] = Form([]), modo: list[str] = Form([]), cantidad: list[str] = Form([]),
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    if motivo not in MOTIVOS_AJUSTE_INVENTARIO:
        flash(request, "Elige el motivo del ajuste.", "error")
        return RedirectResponse(BASE + "/nuevo", status_code=303)
    if motivo == "Otro" and not observacion.strip():
        flash(request, "Con el motivo «Otro», escribe en la observación qué pasó.", "error")
        return RedirectResponse(BASE + "/nuevo", status_code=303)
    lineas = [
        {"producto_id": pid, "modo": m, "cantidad": c}
        for pid, m, c in zip(producto_id, modo, cantidad) if (pid or "").strip()
    ]
    if not lineas:
        flash(request, "Agrega al menos un producto al ajuste.", "error")
        return RedirectResponse(BASE + "/nuevo", status_code=303)
    ajuste, errores = aplicar_ajuste(db, lineas, motivo, observacion, "manual", usuario["nombre_completo"])
    if errores:
        db.rollback()
        flash(request, "No se aplicó el ajuste: " + " ".join(errores), "error")
        return RedirectResponse(BASE + "/nuevo", status_code=303)
    db.commit()
    flash(request, f"Ajuste {ajuste.numero} aplicado a {len(ajuste.detalles)} producto(s).", "success")
    return RedirectResponse(f"{BASE}/{ajuste.id}", status_code=303)


# ---------------------------------------------------------------------------
# Importar desde Excel
# ---------------------------------------------------------------------------
@router.get(BASE + "/importar")
def ajustes_importar_form(request: Request, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    categorias = [c.nombre for c in db.query(Categoria).order_by(Categoria.nombre).all()]
    return templates.TemplateResponse("ajustes_inventario/importar.html", {
        "request": request, "usuario": usuario, "motivos": MOTIVOS_AJUSTE_INVENTARIO, "categorias": categorias,
    })


@router.get(BASE + "/plantilla")
def ajustes_plantilla(categoria: str = "", db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    productos = _productos_activos(db)
    if categoria:
        productos = [p for p in productos if (p.categoria or "") == categoria]
    wb = Workbook()
    hoja = wb.active
    hoja.title = "Ajuste"
    encabezados = ["Código", "Nombre", "Categoría", "Existencia en sistema", "Cantidad"]
    hoja.append(encabezados)
    relleno = PatternFill("solid", fgColor="1F4E78")
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")
    for p in productos:
        hoja.append([p.codigo, p.nombre, p.categoria or "", p.existencia or 0, None])
    for col, ancho in zip("ABCDE", (16, 42, 18, 22, 14)):
        hoja.column_dimensions[col].width = ancho
    validacion = DataValidation(type="whole", allow_blank=True, operator="between", formula1="-1000000", formula2="1000000",
                                error="Escribe un número entero.", errorTitle="Cantidad inválida")
    hoja.add_data_validation(validacion)
    validacion.add(f"E2:E{max(len(productos) + 1, 500)}")
    hoja.freeze_panes = "A2"

    instrucciones = wb.create_sheet("Instrucciones")
    for fila in (
        ("Cómo llenar esta plantilla",),
        ("1.", "Escribe la cantidad solo en los productos que quieres ajustar; deja en blanco los demás (se ignoran)."),
        ("2.", "Al subir el archivo eliges qué significa la columna Cantidad:"),
        ("", "   • Existencia final (conteo): la cantidad que contaste. El sistema calcula cuánto sube o baja."),
        ("", "   • Diferencia (+/-): cuánto sumar (ej. 3) o restar (ej. -2) a lo que ya hay."),
        ("3.", "No cambies la columna Código. Nombre, Categoría y Existencia en sistema son solo de referencia."),
        ("4.", "Antes de aplicar verás una vista previa con lo que va a cambiar y su valor en lempiras."),
    ):
        instrucciones.append(list(fila))
    instrucciones["A1"].font = Font(bold=True, size=13)
    instrucciones.column_dimensions["A"].width = 6
    instrucciones.column_dimensions["B"].width = 100

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return StreamingResponse(
        salida, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="plantilla_ajuste_inventario.xlsx"'},
    )


def _texto_codigo(valor) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _leer_filas_excel(contenido: bytes):
    wb = load_workbook(io.BytesIO(contenido), data_only=True)
    hoja = wb["Ajuste"] if "Ajuste" in wb.sheetnames else wb.worksheets[0]
    filas = list(hoja.iter_rows(values_only=True))
    if not filas:
        return []
    encabezado = [str(c or "").strip().lower() for c in filas[0]]
    try:
        col_codigo = next(i for i, c in enumerate(encabezado) if c.startswith("c") and "digo" in c)
        col_cantidad = next(i for i, c in enumerate(encabezado) if c.startswith("cantidad"))
    except StopIteration:
        raise ValueError("El Excel debe tener las columnas «Código» y «Cantidad» en la primera fila (usa la plantilla).")
    resultado = []
    for numero_fila, fila in enumerate(filas[1:], start=2):
        codigo = _texto_codigo(fila[col_codigo] if col_codigo < len(fila) else None)
        cantidad = fila[col_cantidad] if col_cantidad < len(fila) else None
        if cantidad is None or str(cantidad).strip() == "":
            continue  # sin cantidad = no se ajusta (así viene la plantilla)
        resultado.append({"fila": numero_fila, "codigo": codigo, "cantidad": cantidad})
    return resultado


@router.post(BASE + "/importar")
async def ajustes_importar_preview(
    request: Request, archivo: UploadFile = File(...), modo_excel: str = Form("final"),
    motivo: str = Form(""), observacion: str = Form(""),
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    destino = BASE + "/importar"
    if motivo not in MOTIVOS_AJUSTE_INVENTARIO:
        flash(request, "Elige el motivo del ajuste.", "error")
        return RedirectResponse(destino, status_code=303)
    if motivo == "Otro" and not observacion.strip():
        flash(request, "Con el motivo «Otro», escribe en la observación qué pasó.", "error")
        return RedirectResponse(destino, status_code=303)
    if modo_excel not in ("final", "diferencia"):
        modo_excel = "final"
    if not (archivo.filename or "").lower().endswith((".xlsx", ".xlsm")):
        flash(request, "El archivo debe ser un Excel (.xlsx). Descarga la plantilla e inténtalo de nuevo.", "error")
        return RedirectResponse(destino, status_code=303)
    try:
        filas_excel = _leer_filas_excel(await archivo.read())
    except ValueError as e:
        flash(request, str(e), "error")
        return RedirectResponse(destino, status_code=303)
    except Exception:
        flash(request, "No se pudo leer el archivo. Asegúrate de que sea un Excel válido (.xlsx).", "error")
        return RedirectResponse(destino, status_code=303)

    productos_por_codigo = {p.codigo.strip().lower(): p for p in db.query(Producto).all() if p.codigo}
    vistos = {}
    filas_preview, lineas_validas = [], []
    for f in filas_excel:
        info = {"fila": f["fila"], "codigo": f["codigo"], "nombre": "", "estado": "error", "mensaje": "",
                "anterior": None, "nueva": None, "diferencia": None, "valor": None}
        producto = productos_por_codigo.get(f["codigo"].lower())
        if not f["codigo"]:
            info["mensaje"] = "Falta el código."
        elif producto is None:
            info["mensaje"] = "No existe un producto con ese código."
        else:
            info["nombre"] = producto.nombre
            try:
                valor_cantidad = float(str(f["cantidad"]).strip())
                if not valor_cantidad.is_integer():
                    raise ValueError
                cantidad = int(valor_cantidad)
            except (TypeError, ValueError):
                cantidad = None
            if cantidad is None:
                info["mensaje"] = "La cantidad debe ser un número entero."
            elif producto.id in vistos:
                info["mensaje"] = f"Código repetido (también en la fila {vistos[producto.id]})."
            elif modo_excel == "final" and cantidad < 0:
                info["mensaje"] = "La existencia final no puede ser negativa."
            else:
                vistos[producto.id] = f["fila"]
                if modo_excel == "final":
                    linea = {"producto_id": producto.id, "modo": "final", "cantidad": cantidad}
                else:
                    linea = {"producto_id": producto.id, "modo": "sumar" if cantidad >= 0 else "restar", "cantidad": abs(cantidad)}
                calculadas, errores = calcular_lineas(db, [linea])
                if errores:
                    info["mensaje"] = errores[0]
                else:
                    c = calculadas[0]
                    info.update({"anterior": c["existencia_anterior"], "nueva": c["existencia_nueva"],
                                 "diferencia": c["diferencia"], "valor": c["valor"]})
                    if c["diferencia"] == 0:
                        info["estado"], info["mensaje"] = "sin_cambio", "Ya tiene esa existencia."
                    else:
                        info["estado"] = "ok"
                        lineas_validas.append(linea)
        filas_preview.append(info)

    ok = [f for f in filas_preview if f["estado"] == "ok"]
    return templates.TemplateResponse("ajustes_inventario/importar_preview.html", {
        "request": request, "usuario": usuario, "filas": filas_preview,
        "total_ok": len(ok), "total_sin_cambio": sum(1 for f in filas_preview if f["estado"] == "sin_cambio"),
        "total_error": sum(1 for f in filas_preview if f["estado"] == "error"),
        "unidades_entrada": sum(f["diferencia"] for f in ok if f["diferencia"] > 0),
        "unidades_salida": sum(-f["diferencia"] for f in ok if f["diferencia"] < 0),
        "valor_entrada": sum((f["valor"] for f in ok if f["diferencia"] > 0), Decimal("0.00")),
        "valor_salida": sum((-f["valor"] for f in ok if f["diferencia"] < 0), Decimal("0.00")),
        "lineas_json": json.dumps(lineas_validas), "motivo": motivo, "observacion": observacion,
        "modo_excel": modo_excel, "nombre_archivo": archivo.filename,
    })


@router.post(BASE + "/importar/confirmar")
def ajustes_importar_confirmar(
    request: Request, lineas_json: str = Form(...), motivo: str = Form(""), observacion: str = Form(""),
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    try:
        lineas = json.loads(lineas_json)
        assert isinstance(lineas, list)
    except (ValueError, TypeError, AssertionError):
        flash(request, "La vista previa no se pudo leer. Vuelve a subir el archivo.", "error")
        return RedirectResponse(BASE + "/importar", status_code=303)
    if motivo not in MOTIVOS_AJUSTE_INVENTARIO:
        flash(request, "Motivo inválido. Vuelve a subir el archivo.", "error")
        return RedirectResponse(BASE + "/importar", status_code=303)
    # Se recalcula todo contra la existencia ACTUAL (si hubo una venta entre
    # la vista previa y la confirmación, "Existencia final" sigue dejando el
    # producto en lo contado y "Diferencia" suma/resta sobre lo que hay hoy).
    ajuste, errores = aplicar_ajuste(db, lineas, motivo, observacion, "excel", usuario["nombre_completo"])
    if errores:
        db.rollback()
        flash(request, "No se aplicó el ajuste: " + " ".join(errores), "error")
        return RedirectResponse(BASE + "/importar", status_code=303)
    db.commit()
    flash(request, f"Ajuste {ajuste.numero} importado: {len(ajuste.detalles)} producto(s) actualizados.", "success")
    return RedirectResponse(f"{BASE}/{ajuste.id}", status_code=303)


# ---------------------------------------------------------------------------
# Reportes
# ---------------------------------------------------------------------------
@router.get(BASE + "/reporte")
def ajustes_reporte(
    request: Request, desde: str = "", hasta: str = "", motivo: str = "", q: str = "",
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    d, h, inicio, fin = _rango(desde, hasta, por_defecto_mes=True)
    datos = reporte_ajustes(db, inicio, fin, motivo, q)
    return templates.TemplateResponse("ajustes_inventario/reporte.html", {
        "request": request, "usuario": usuario, "datos": datos,
        "desde": d.isoformat() if d else "", "hasta": h.isoformat() if h else "",
        "motivo": motivo, "q": q, "hoy_iso": hoy_local().isoformat(),
        "motivos": MOTIVOS_AJUSTE_INVENTARIO + [MOTIVO_REVERSION_AJUSTE, MOTIVO_AJUSTE_RAPIDO],
        "a_hora_local": a_hora_local,
    })


@router.get(BASE + "/reporte/excel")
def ajustes_reporte_excel(
    desde: str = "", hasta: str = "", motivo: str = "", q: str = "",
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    d, h, inicio, fin = _rango(desde, hasta, por_defecto_mes=True)
    datos = reporte_ajustes(db, inicio, fin, motivo, q)
    wb = Workbook()
    negrita = Font(bold=True)
    relleno = PatternFill("solid", fgColor="1F4E78")

    def encabezado(hoja, columnas):
        hoja.append(columnas)
        for celda in hoja[hoja.max_row]:
            celda.font = Font(bold=True, color="FFFFFF")
            celda.fill = relleno

    def grupo(hoja, titulo, filas):
        hoja.append([])
        hoja.append([titulo])
        hoja.cell(row=hoja.max_row, column=1).font = negrita
        encabezado(hoja, [titulo.split(" por ")[-1].capitalize(), "Unidades que subieron", "Unidades que bajaron",
                          "Valor que subió (L)", "Valor que bajó (L)", "Neto (L)"])
        for g in filas:
            hoja.append([g["clave"], g["unidades_entrada"], g["unidades_salida"],
                         float(g["valor_entrada"]), float(g["valor_salida"]), float(g["valor_neto"])])

    resumen = wb.active
    resumen.title = "Resumen"
    rango_txt = f"{d.strftime('%d/%m/%Y') if d else 'inicio'} al {h.strftime('%d/%m/%Y') if h else 'hoy'}"
    resumen.append([f"Reporte de ajustes de inventario — {rango_txt}"])
    resumen["A1"].font = Font(bold=True, size=13)
    t = datos["total"]
    resumen.append(["Unidades que subieron", t["unidades_entrada"]])
    resumen.append(["Unidades que bajaron", t["unidades_salida"]])
    resumen.append(["Valor que subió (L)", float(t["valor_entrada"])])
    resumen.append(["Valor que bajó (L)", float(t["valor_salida"])])
    resumen.append(["Neto (L)", float(t["valor_neto"])])
    grupo(resumen, "Totales por motivo", datos["por_motivo"])
    grupo(resumen, "Totales por categoría", datos["por_categoria"])
    grupo(resumen, "Totales por mes", datos["por_mes"])
    if datos["hay_aproximados"]:
        resumen.append([])
        resumen.append(["Nota: los ajustes rápidos anteriores al módulo se valoran con el costo actual del producto (aproximado)."])

    hoja_prod = wb.create_sheet("Por producto")
    encabezado(hoja_prod, ["Producto", "Unidades que subieron", "Unidades que bajaron", "Neto unidades",
                           "Valor que subió (L)", "Valor que bajó (L)", "Neto (L)"])
    for g in datos["por_producto"]:
        hoja_prod.append([g["clave"], g["unidades_entrada"], g["unidades_salida"], g["unidades_neto"],
                          float(g["valor_entrada"]), float(g["valor_salida"]), float(g["valor_neto"])])

    detalle = wb.create_sheet("Detalle")
    encabezado(detalle, ["Fecha", "Ajuste", "Motivo", "Código", "Producto", "Categoría", "Diferencia",
                         "Costo unitario (L)", "Valor (L)", "Usuario", "Nota", "Aproximado"])
    for f in datos["filas"]:
        detalle.append([
            a_hora_local(f["fecha"]).strftime("%d/%m/%Y %I:%M %p"), f["ajuste"].numero if f["ajuste"] else "",
            f["motivo"], f["codigo"], f["nombre"], f["categoria"], f["diferencia"],
            float(f["costo_unitario"]), float(f["valor"]), f["usuario"], f["detalle_motivo"],
            "Sí" if f["aproximado"] else "",
        ])
    for hoja in (resumen, hoja_prod, detalle):
        hoja.column_dimensions["A"].width = 34
        for col in "BCDEFGHIJKL":
            hoja.column_dimensions[col].width = 18
    detalle.column_dimensions["E"].width = 40

    salida = io.BytesIO()
    wb.save(salida)
    salida.seek(0)
    return StreamingResponse(
        salida, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="reporte_ajustes_inventario.xlsx"'},
    )


# ---------------------------------------------------------------------------
# Detalle y reversión
# ---------------------------------------------------------------------------
@router.get(BASE + "/{ajuste_id}")
def ajustes_detalle(ajuste_id: int, request: Request, db: Session = Depends(get_db), usuario=Depends(solo_admin)):
    ajuste = db.get(AjusteInventario, ajuste_id)
    if not ajuste:
        flash(request, "Ajuste no encontrado.", "error")
        return RedirectResponse(BASE, status_code=303)
    return templates.TemplateResponse("ajustes_inventario/detalle.html", {
        "request": request, "usuario": usuario, "ajuste": ajuste,
        "revierte_a": db.get(AjusteInventario, ajuste.revierte_a_id) if ajuste.revierte_a_id else None,
        "revertido_por": db.get(AjusteInventario, ajuste.revertido_por_id) if ajuste.revertido_por_id else None,
        "a_hora_local": a_hora_local,
    })


@router.post(BASE + "/{ajuste_id}/revertir")
def ajustes_revertir(
    ajuste_id: int, request: Request, nota: str = Form(""),
    db: Session = Depends(get_db), usuario=Depends(solo_admin),
):
    ajuste = db.get(AjusteInventario, ajuste_id)
    if not ajuste:
        flash(request, "Ajuste no encontrado.", "error")
        return RedirectResponse(BASE, status_code=303)
    reversion, errores = revertir_ajuste(db, ajuste, usuario["nombre_completo"], nota)
    if errores:
        db.rollback()
        flash(request, "No se pudo revertir: " + " ".join(errores), "error")
        return RedirectResponse(f"{BASE}/{ajuste_id}", status_code=303)
    db.commit()
    flash(request, f"Ajuste {ajuste.numero} revertido con el ajuste {reversion.numero}.", "success")
    return RedirectResponse(f"{BASE}/{reversion.id}", status_code=303)
